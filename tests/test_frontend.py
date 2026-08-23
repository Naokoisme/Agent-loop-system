from __future__ import annotations

import copy
import json
import os
import tempfile
import threading
import unittest
from datetime import datetime, timedelta
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch
from urllib.error import HTTPError
from urllib.parse import quote, urlencode
from urllib.request import Request, urlopen

from frontend.server import (
    AppPaths,
    BATCH_STATE_FILE,
    CaseMapRepository,
    CaseTestManager,
    DefectRepository,
    FrontendHTTPServer,
    HistoryStore,
    JobManager,
    RequestHandler,
    TestHistoryStore as CaseRunHistoryStore,
    ThreadingHTTPServer,
    WebApplication,
    _save_system_config,
    main as frontend_main,
    make_handler,
)


def _external_ledger_record(
    case_id: str,
    sheet: str,
    target: str,
) -> dict[str, object]:
    return {
        "case_id": case_id,
        "sheet": sheet,
        "target": target,
        "last_verified": "2026-08-17",
        "evidence_root": "D:/external-evidence",
        "evidence_paths": [f"tests/{case_id}/result.json"],
    }


class FrontendDataTest(unittest.TestCase):
    def _restore_env(self) -> None:
        os.environ.clear()
        os.environ.update(self._env_backup)

    def test_system_config_persists_layout_paths_relatively(self) -> None:
        layout_root = self.paths.root.parent
        simulator_root = layout_root / "workspaces" / "firmware" / "6202_W5230"
        with patch.dict(
            os.environ,
            {
                "AGENT_LOOP_ROOT": str(self.paths.root),
                "AGENT_LOOP_LAYOUT_ROOT": "..",
            },
            clear=False,
        ):
            _save_system_config(
                self.paths,
                {
                    "simulator_6202": {
                        "source_root": str(simulator_root),
                        "build_directory": str(simulator_root / "build"),
                        "artifact_path": str(simulator_root / "bin" / "main.exe"),
                    }
                },
            )

        values = {
            key: value
            for line in (self.paths.root / ".env").read_text(encoding="utf-8").splitlines()
            if line and "=" in line
            for key, value in [line.split("=", 1)]
        }
        self.assertEqual(
            values["W30_6202_SIMULATOR_SOURCE_ROOT"],
            os.path.relpath(simulator_root.resolve(), self.paths.root),
        )
        self.assertEqual(
            values["W30_6202_SIMULATOR_BUILD_DIRECTORY"],
            os.path.relpath((simulator_root / "build").resolve(), self.paths.root),
        )

    def setUp(self) -> None:
        self._env_backup = dict(os.environ)
        self.addCleanup(self._restore_env)
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        root = Path(self.temporary.name)
        self.paths = AppPaths.from_root(root)
        for path in (
            self.paths.frontend,
            self.paths.defects / "100",
            self.paths.defect_images,
            self.paths.evidence / "100",
            self.paths.case_map / "620C_simulator_case_map",
            self.paths.case_map / "6202_case_map",
            self.paths.case_map / "6202_simulator_case_map",
        ):
            path.mkdir(parents=True, exist_ok=True)
        (self.paths.frontend / "index.html").write_text("<h1>正常</h1>", encoding="utf-8")

        self._write_defect(
            "100",
            title="计算器显示错误",
            description="复现描述：点击等号后显示异常",
            status="处理中",
        )
        (self.paths.case_map / "620C_simulator_case_map" / "计算器.json").write_text(
            json.dumps(
                [
                    {
                        "case_id": "CALC_001",
                        "sheet": "计算器",
                        "priority": "P0",
                        "precondition_text": "已进入计算器",
                        "steps_text": "点击等号",
                        "expected_text": "显示正确",
                        "verification_points": ["结果区域显示正确数值"],
                        "setup": ["srv_quick_cmd send TOP5STEP:ENTER_PAGE:CALCULATOR,0;"],
                        "actions": ["srv_quick_cmd send TOP5STEP:TP_CLICK:10,20,1;"],
                        "collect": ["srv_quick_cmd send TOP5STEP:GUI_TREE:1;"],
                        "unable": False,
                        "mapping_status": "PROMOTED",
                    },
                    {
                        "case_id": "CALC_002",
                        "sheet": "计算器",
                        "priority": "P1",
                        "precondition_text": "连接手机 App",
                        "steps_text": "从 App 下发数据",
                        "expected_text": "显示同步结果",
                        "setup": [],
                        "actions": [],
                        "collect": ["srv_quick_cmd send TOP5STEP:GUI_TREE:1;"],
                        "unable": True,
                        "note": "依赖手机 App",
                    }
                ],
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        (self.paths.case_map / "6202_case_map" / "计算器.json").write_text(
            json.dumps(
                {
                    "profile": "6202_W5230",
                    "sheet": "计算器",
                    "cases": [
                    {
                        "case_id": "CALC_001",
                        "sheet": "计算器",
                        "priority": "P0",
                        "precondition_text": "真机已进入计算器",
                        "steps_text": "在真机点击等号",
                        "expected_text": "真机显示正确",
                        "verification_points": ["真机结果区域显示正确数值"],
                        "setup": ["srv_quick_cmd send TOP5STEP:ENTER_PAGE:CALCULATOR,0;"],
                        "actions": ["srv_quick_cmd send TOP5STEP:TP_CLICK:10,20,1;"],
                        "collect": ["srv_quick_cmd send TOP5STEP:GUI_TREE:1;"],
                        "unable": False,
                    }
                    ],
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        (self.paths.case_map / "6202_simulator_case_map" / "计算器.json").write_text(
            json.dumps(
                {
                    "profile": "6202_W5230_SIMULATOR",
                    "sheet": "计算器",
                    "cases": [{
                        "case_id": "CALC_001",
                        "sheet": "计算器",
                        "priority": "P0",
                        "precondition_text": "6202 模拟器已启动",
                        "steps_text": "在 6202 模拟器点击等号",
                        "expected_text": "6202 模拟器显示正确",
                        "verification_points": ["截图中结果区域显示正确数值"],
                        "setup": ["srv_quick_cmd send TOP5STEP:ENTER_PAGE:CALCULATOR,0;"],
                        "actions": ["srv_quick_cmd send TOP5STEP:TP_CLICK:10,20,1;"],
                        "collect": ["srv_quick_cmd send TOP5STEP:HOST_SCREENSHOT:1;"],
                        "unable": False,
                    }],
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        for directory, target, records in (
            ("620C_simulator_case_map", "620C_W6830", [("CALC_001", "计算器")]),
            ("6202_case_map", "6202_W5230", [("CALC_001", "计算器")]),
            (
                "6202_simulator_case_map",
                "6202_W5230_SIMULATOR",
                [("CALC_001", "计算器")],
            ),
        ):
            ledger = self.paths.case_map / directory / "external_execution_history.jsonl"
            ledger.write_text(
                "".join(
                    json.dumps(
                        _external_ledger_record(case_id, sheet, target),
                        ensure_ascii=False,
                    ) + "\n"
                    for case_id, sheet in records
                ),
                encoding="utf-8",
            )
        (self.paths.evidence / "100" / "before.bmp").write_bytes(b"BM-before")
        (self.paths.evidence / "100" / "after.bmp").write_bytes(b"BM-after")
        self.history = HistoryStore(self.paths)
        self.defects = DefectRepository(self.paths, self.history)
        self.test_history = CaseRunHistoryStore(self.paths)
        self.cases = CaseMapRepository(self.paths, self.test_history)

    def test_frontend_server_rejects_a_second_process_on_the_same_port(self) -> None:
        application = WebApplication(self.paths)
        first = FrontendHTTPServer(
            ("127.0.0.1", 0),
            make_handler(application),
        )
        self.addCleanup(first.server_close)

        with self.assertRaises(OSError):
            FrontendHTTPServer(
                ("127.0.0.1", first.server_address[1]),
                make_handler(application),
            )

    def test_frontend_main_uses_the_runtime_app_root(self) -> None:
        runtime_root = Path(self.temporary.name) / "portable-app"
        fake_server = SimpleNamespace(
            serve_forever=lambda: None,
            server_close=lambda: None,
        )

        with (
            patch("frontend.server.resolve_app_root", return_value=runtime_root),
            patch("frontend.server.WebApplication") as application_class,
            patch("frontend.server.FrontendHTTPServer", return_value=fake_server),
        ):
            self.assertEqual(frontend_main(["--host", "127.0.0.1", "--port", "0"]), 0)

        paths = application_class.call_args.args[0]
        self.assertEqual(paths.root, runtime_root.resolve())
        self.assertEqual(paths.frontend, runtime_root.resolve() / "frontend")

    def _write_defect(
        self,
        number: str,
        *,
        title: str,
        description: str = "普通缺陷",
        status: str = "待处理",
        attachments: list[dict[str, object]] | None = None,
    ) -> None:
        defect_dir = self.paths.defects / number
        defect_dir.mkdir(parents=True, exist_ok=True)
        (defect_dir / "defect.json").write_text(
            json.dumps(
                {
                    "number": number,
                    "title": title,
                    "description": description,
                    "status": status,
                    "attachments": attachments or [],
                    "source_analysis": {
                        "matches": [
                            {
                                "path": "app/calculator.c",
                                "line_start": 10,
                                "line_end": 12,
                                "snippet": "old();",
                            }
                        ]
                    },
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )

    def _create_history(self, verdict: str = "CANNOT_VERIFY", defect: str = "100") -> str:
        return self.history.create(
            defect=defect,
            job={
                "started_at": "2026-08-08T10:00:00+08:00",
                "finished_at": "2026-08-08T10:01:00+08:00",
                "execution_mode": "agent_generated",
                "progress": {"nodes": {"validate": "pass", "test": "fail"}},
            },
            result={
                "verdict": verdict,
                "attempts": 1,
                "patch_retained": False,
                "patch": {
                    "file_path": "app/calculator.c",
                    "before": "old",
                    "after": "new",
                    "reason": "修复显示逻辑",
                    "test_commands": [":GUI_TREE:1"],
                },
                "test_output": {
                    "results": [
                        {
                            "verdict": verdict,
                            "reason": "模拟器未返回完整数据",
                            "test_commands": [":GUI_TREE:1", ":GET_CURRENT_WIN_ID"],
                            "terminal_json": [{"ok": False}],
                        }
                    ]
                },
            },
            stdout="done",
            stderr="",
        )

    def _write_history_record(
        self,
        run_id: str,
        run: dict[str, object],
        *,
        test_result: dict[str, object] | None = None,
    ) -> None:
        run_dir = self.paths.history / "100" / run_id
        run_dir.mkdir(parents=True)
        (run_dir / "run.json").write_text(
            json.dumps(run, ensure_ascii=False),
            encoding="utf-8",
        )
        if test_result is not None:
            (run_dir / "test_result.json").write_text(
                json.dumps(test_result, ensure_ascii=False),
                encoding="utf-8",
            )

    def _server(self) -> tuple[WebApplication, str]:
        application = WebApplication(self.paths)
        server = ThreadingHTTPServer(("127.0.0.1", 0), make_handler(application))
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        self.addCleanup(server.server_close)
        self.addCleanup(server.shutdown)
        return application, f"http://127.0.0.1:{server.server_port}"

    def _clear_case_candidate(
        self,
        *,
        project_dir: str,
        sheet: str = "计算器",
        case_id: str = "CALC_001",
    ) -> Path:
        path = self.paths.case_map / project_dir / f"{sheet}.json"
        raw = json.loads(path.read_text(encoding="utf-8"))
        entries = raw["cases"] if isinstance(raw, dict) else raw
        case = next(item for item in entries if item["case_id"] == case_id)
        for field in ("setup", "actions", "collect", "verification_points"):
            case[field] = []
        case["note"] = ""
        case.pop("mapping_status", None)
        path.write_text(json.dumps(raw, ensure_ascii=False, indent=2), encoding="utf-8")
        return path

    @staticmethod
    def _agent_exploration_source(
        *,
        project: str = "6202_W5230",
        sheet: str = "计算器",
        case_id: str = "CALC_001",
    ) -> dict[str, object]:
        points = ["执行前初始画面", "点击后结果画面"]
        return {
            "id": "20260821T120000000000",
            "project": project,
            "sheet": sheet,
            "case_id": case_id,
            "execution_mode": "agent_exploration",
            "verdict": "PASS",
            "reason": "真实业务动作已执行，截图符合预期",
            "verification_points": points,
            "planned_commands": {
                "setup": [],
                "action": [":TP_CLICK:10,20,0"],
                "collect": [],
            },
            "command_trace": [
                {
                    "index": 1,
                    "phase": "exploration",
                    "source": "runner",
                    "kind": "capture",
                    "command": ":HOST_SCREENSHOT:AGENT_EXPLORATION",
                    "command_name": "HOST_SCREENSHOT",
                    "ok": True,
                    "checkpoint_index": 1,
                    "checkpoint_label": points[0],
                },
                {
                    "index": 2,
                    "phase": "action",
                    "source": "agent",
                    "kind": "device",
                    "command": ":TP_CLICK:10,20,0",
                    "command_name": "TP_CLICK",
                    "ok": True,
                },
                {
                    "index": 3,
                    "phase": "exploration",
                    "source": "runner",
                    "kind": "capture",
                    "command": ":HOST_SCREENSHOT:AGENT_EXPLORATION",
                    "command_name": "HOST_SCREENSHOT",
                    "ok": True,
                    "checkpoint_index": 2,
                    "checkpoint_label": points[1],
                },
            ],
            "evidence_contract": {
                "status": "COMPLETE",
                "complete": True,
                "issues": [],
                "required_screenshots": 2,
                "captured_screenshots": 2,
                "business_action_count": 1,
                "planned_action_count": 1,
                "attempted_action_count": 1,
            },
            "screenshots": [{"file": "one.bmp"}, {"file": "two.bmp"}],
            "screenshot_urls": [{"file": "one.bmp"}, {"file": "two.bmp"}],
            "skipped": False,
            "aborted": False,
            "setup_errors": [],
            "action_errors": [],
            "collect_errors": [],
            "exploration_trace": {
                "steps": [{"step": 0, "window_name": "CALCULATOR"}],
            },
        }

    @staticmethod
    def _write_valid_hardware_bmp(path: Path) -> None:
        import struct

        width, height = 410, 502
        pixel_offset = 54
        row_stride = ((width * 3) + 3) & ~3
        file_size = pixel_offset + row_stride * height
        data = bytearray(file_size)
        data[:2] = b"BM"
        struct.pack_into("<I", data, 2, file_size)
        struct.pack_into("<I", data, 10, pixel_offset)
        struct.pack_into("<I", data, 14, 40)
        struct.pack_into("<ii", data, 18, width, -height)
        struct.pack_into("<HH", data, 26, 1, 24)
        struct.pack_into("<I", data, 30, 0)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(data)

    def _candidate_replay_result(
        self,
        context: dict[str, object],
        *,
        job_id: str,
        verdict: str = "PASS",
    ) -> dict[str, object]:
        candidate = context["candidate_fields"]
        context["job_id"] = job_id
        evidence_dir = self.paths.runtime_jobs / job_id / "single"
        first = evidence_dir / "screenshot-01.bmp"
        second = evidence_dir / "screenshot-02.bmp"
        self._write_valid_hardware_bmp(first)
        self._write_valid_hardware_bmp(second)
        points = candidate["verification_points"]
        return {
            "schema_version": 3,
            "case_id": "CALC_001",
            "sheet": "计算器",
            "execution_mode": "candidate_mapping",
            "verdict": verdict,
            "reason": "候选复跑形成确定产品结论",
            "execution_status": "OK",
            "verification_points": points,
            "planned_commands": {
                "setup": candidate["setup"],
                "action": candidate["actions"],
                "collect": candidate["collect"],
            },
            "command_trace": [
                {
                    "index": 1,
                    "phase": "setup",
                    "source": "case",
                    "kind": "screenshot",
                    "wire": candidate["setup"][1],
                    "command": candidate["setup"][1],
                    "command_name": "HOST_SCREENSHOT",
                    "planned_index": 2,
                    "checkpoint_index": 1,
                    "checkpoint_label": points[0],
                    "ok": True,
                },
                {
                    "index": 2,
                    "phase": "action",
                    "source": "case",
                    "kind": "device",
                    "wire": candidate["actions"][0],
                    "command": candidate["actions"][0],
                    "command_name": "TP_CLICK",
                    "planned_index": 1,
                    "ok": True,
                },
                {
                    "index": 3,
                    "phase": "action",
                    "source": "case",
                    "kind": "screenshot",
                    "wire": candidate["actions"][1],
                    "command": candidate["actions"][1],
                    "command_name": "HOST_SCREENSHOT",
                    "planned_index": 2,
                    "checkpoint_index": 2,
                    "checkpoint_label": points[1],
                    "ok": True,
                },
            ],
            "evidence_contract": {
                "status": "COMPLETE",
                "complete": True,
                "issues": [],
                "required_screenshots": 2,
                "captured_screenshots": 2,
                "planned_action_count": 2,
                "attempted_action_count": 2,
                "business_action_count": 1,
            },
            "screenshots": [
                {
                    "path": str(first),
                    "label": points[0],
                    "captured_at": "2026-08-21T12:01:00+08:00",
                    "trace_index": 1,
                },
                {
                    "path": str(second),
                    "label": points[1],
                    "captured_at": "2026-08-21T12:02:00+08:00",
                    "trace_index": 3,
                },
            ],
            "provenance": {
                "target": "hardware",
                "case_map_profile": "6202_W5230",
                "project": "6202_W5230",
                "artifact_path": "",
                "artifact_sha256": "",
            },
            "skipped": False,
            "aborted": False,
            "setup_errors": [],
            "action_errors": [],
            "collect_errors": [],
        }


    @staticmethod
    def _put_json(url: str, payload: dict[str, object]):
        request = Request(
            url,
            data=json.dumps(payload).encode("utf-8"),
            headers={"Content-Type": "application/json"},
            method="PUT",
        )
        return urlopen(request, timeout=3)

    @staticmethod
    def _post_json(url: str, payload: dict[str, object]):
        request = Request(
            url,
            data=json.dumps(payload).encode("utf-8"),
            headers={"Content-Type": "application/json"},
            method="POST",
        )
        return urlopen(request, timeout=3)

    @staticmethod
    def _delete_json(url: str):
        return urlopen(Request(url, method="DELETE"), timeout=3)

    def test_defect_detail_keeps_read_only_source_and_case_compatibility(self) -> None:
        detail = self.defects.get("100")
        self.assertIsNotNone(detail)
        self.assertEqual(detail["source_files"][0]["path"], "app/calculator.c")
        self.assertIn("1 处命中", detail["source_files"][0]["label"])
        self.assertEqual(detail["case_sheets"][0]["name"], "计算器")
        self.assertEqual(self.defects.cases("计算器")[0]["case_id"], "CALC_001")

    def test_case_map_list_filters_and_keeps_full_detail(self) -> None:
        payload = self.cases.list()
        self.assertEqual(payload["summary"], {
            "all": 2,
            "unexplored": 1,
            "externally_explored": 1,
            "explored_unsolidified": 0,
            "solidified": 1,
        })
        self.assertEqual(payload["batch_summary"], {
            "untested": 2,
            "fail": 0,
            "cannot_verify": 0,
            "error": 0,
            "pass": 0,
        })
        self.assertEqual(payload["module_counts"], {"计算器": 2})
        self.assertEqual(payload["catalog_total"], 2)
        self.assertEqual(payload["verdict_summary"]["PENDING"], 2)
        queried = self.cases.list(query="显示正确")
        self.assertEqual(queried["summary"]["all"], 1)
        self.assertEqual([row["case_id"] for row in queried["items"]], ["CALC_001"])
        self.assertEqual(
            [row["case_id"] for row in self.cases.list(modules={"计算器"})["items"]],
            ["CALC_001", "CALC_002"],
        )
        missing_module = self.cases.list(modules={"不存在的模块"})
        self.assertEqual(missing_module["items"], [])
        self.assertEqual(missing_module["summary"]["all"], 0)
        self.assertEqual(missing_module["module_counts"], {"计算器": 2})
        self.assertEqual([row["case_id"] for row in self.cases.list(state_filter="solidified")["items"]], ["CALC_001"])
        self.assertEqual([row["case_id"] for row in self.cases.list(state_filter="externally_explored")["items"]], ["CALC_001"])
        self.assertEqual([row["case_id"] for row in self.cases.list(state_filter="unexplored")["items"]], ["CALC_002"])
        self.assertEqual(self.cases.list(state_filter="explored_unsolidified")["items"], [])
        detail = self.cases.get("计算器", "CALC_001")
        self.assertEqual(detail["precondition_text"], "已进入计算器")
        self.assertEqual(detail["actions"], ["srv_quick_cmd send TOP5STEP:TP_CLICK:10,20,1;"])
        self.assertEqual(detail["verification_points"], ["结果区域显示正确数值"])
        self.assertTrue(detail["is_promoted"])
        self.assertEqual(detail["mapping_status"], "PROMOTED")
        with self.assertRaisesRegex(ValueError, "state 参数不合法"):
            self.cases.list(state_filter="unknown")
        for legacy_filter in ("untested", "pass", "fail", "cannot_verify", "error"):
            with self.subTest(legacy_filter=legacy_filter):
                with self.assertRaisesRegex(ValueError, "state 参数不合法"):
                    self.cases.list(state_filter=legacy_filter)

    def test_case_map_project_selects_matching_map_and_execution_target(self) -> None:
        simulator = self.cases.list(project="620C_W6830")
        hardware = self.cases.list(project="6202_W5230")
        simulator_6202 = self.cases.list(project="6202_W5230_SIMULATOR")

        self.assertEqual(simulator["project_label"], "620C W6830")
        self.assertEqual(simulator["execution_target"], "simulator")
        self.assertEqual([item["case_id"] for item in simulator["items"]], ["CALC_001", "CALC_002"])
        self.assertEqual(hardware["project_label"], "6202 W5230")
        self.assertEqual(hardware["execution_target"], "hardware")
        self.assertEqual(hardware["execution_target_label"], "真机")
        self.assertEqual([item["case_id"] for item in hardware["items"]], ["CALC_001"])
        self.assertEqual(hardware["items"][0]["steps_text"], "在真机点击等号")
        self.assertEqual(simulator_6202["execution_target"], "simulator")
        self.assertEqual(simulator_6202["execution_target_label"], "模拟器")
        self.assertEqual(
            simulator_6202["items"][0]["steps_text"],
            "在 6202 模拟器点击等号",
        )
        self.assertEqual(len(hardware["projects"]), 3)
        with self.assertRaisesRegex(ValueError, "测试项目不存在"):
            self.cases.list(project="unknown")

    def test_internal_promotion_is_solidified_without_faking_external_history(self) -> None:
        ledger = self.paths.case_map / "6202_case_map" / "external_execution_history.jsonl"
        ledger.write_text("", encoding="utf-8")
        path = self.paths.case_map / "6202_case_map" / "计算器.json"
        raw = json.loads(path.read_text(encoding="utf-8"))
        raw["cases"][0]["mapping_status"] = "PROMOTED"
        path.write_text(json.dumps(raw, ensure_ascii=False), encoding="utf-8")

        detail = self.cases.get("计算器", "CALC_001", project="6202_W5230")

        self.assertTrue(detail["is_promoted"])
        self.assertFalse(detail["external_explored"])
        self.assertEqual(detail["maturity_state"], "solidified")
        self.assertEqual(
            self.cases.list(project="6202_W5230")["summary"],
            {
                "all": 1,
                "unexplored": 0,
                "externally_explored": 0,
                "explored_unsolidified": 0,
                "solidified": 1,
            },
        )

    def test_agent_exploration_candidate_promotes_only_after_formal_replay(self) -> None:
        case_map_path = self._clear_case_candidate(project_dir="6202_case_map")
        ledger = self.paths.case_map / "6202_case_map" / "external_execution_history.jsonl"
        ledger.write_text("", encoding="utf-8")
        ledger_before = ledger.read_bytes()
        source = self._agent_exploration_source()

        context = self.cases.stage_agent_candidate(
            sheet="计算器",
            case_id="CALC_001",
            project="6202_W5230",
            source_history=source,
        )
        candidate = context["candidate_fields"]
        self.assertEqual(
            candidate["setup"],
            [":ENTER_PAGE:CALCULATOR,0", ":HOST_SCREENSHOT:1"],
        )
        self.assertEqual(
            candidate["actions"],
            [":TP_CLICK:10,20,0", ":HOST_SCREENSHOT:2"],
        )
        staged = json.loads(case_map_path.read_text(encoding="utf-8"))["cases"][0]
        self.assertNotIn("mapping_status", staged)
        self.assertEqual(staged["actions"], candidate["actions"])

        job_id = "promotion-job"
        replay_result = self._candidate_replay_result(context, job_id=job_id)

        promotion = self.cases.finalize_agent_candidate(
            context=context,
            result=replay_result,
        )

        self.assertEqual(promotion, {"status": "promoted", "issues": []})
        promoted = json.loads(case_map_path.read_text(encoding="utf-8"))["cases"][0]
        self.assertEqual(promoted["mapping_status"], "PROMOTED")
        self.assertEqual(ledger.read_bytes(), ledger_before)

        self._clear_case_candidate(project_dir="6202_case_map")
        rollback_context = self.cases.stage_agent_candidate(
            sheet="计算器",
            case_id="CALC_001",
            project="6202_W5230",
            source_history=source,
        )
        rollback_context["job_id"] = job_id
        invalid_result = copy.deepcopy(replay_result)
        invalid_result["evidence_contract"]["complete"] = False
        invalid_result["evidence_contract"]["status"] = "ERROR"
        rolled_back = self.cases.finalize_agent_candidate(
            context=rollback_context,
            result=invalid_result,
        )

        self.assertEqual(rolled_back["status"], "rolled_back")
        restored = json.loads(case_map_path.read_text(encoding="utf-8"))["cases"][0]
        self.assertNotIn("mapping_status", restored)
        self.assertEqual(restored["setup"], [])
        self.assertEqual(restored["actions"], [])
        self.assertEqual(restored["verification_points"], [])
        self.assertEqual(ledger.read_bytes(), ledger_before)

    def test_frontend_rejects_case_map_profile_metadata_mismatch(self) -> None:
        path = self.paths.case_map / "6202_simulator_case_map" / "计算器.json"
        payload = json.loads(path.read_text(encoding="utf-8"))
        payload["profile"] = "620C_W6830"
        path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")

        with self.assertRaisesRegex(ValueError, "profile 不匹配"):
            self.cases.list(project="6202_W5230_SIMULATOR")

        payload["profile"] = "6202_W5230_SIMULATOR"
        payload["cases"] = "not-an-array"
        path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "cases 必须是非空数组"):
            self.cases.list(project="6202_W5230_SIMULATOR")

    def test_http_api_rejects_corrupt_620c_case_map_json(self) -> None:
        _, base = self._server()
        path = self.paths.case_map / "620C_simulator_case_map" / "计算器.json"
        path.write_text("{not-json", encoding="utf-8")

        with self.assertRaises(HTTPError) as raised:
            urlopen(base + "/api/tests?project=620C_W6830", timeout=3)
        self.assertEqual(raised.exception.code, 400)
        error = json.loads(raised.exception.read().decode("utf-8"))
        self.assertIn("case_map JSON 无法读取", error["error"])

    def test_only_exact_promoted_status_is_formalized(self) -> None:
        path = self.paths.case_map / "620C_simulator_case_map" / "计算器.json"
        entries = json.loads(path.read_text(encoding="utf-8"))
        template = dict(entries[0])
        entries.extend([
            {**template, "case_id": "CALC_003", "mapping_status": "promoted"},
            {key: value for key, value in {**template, "case_id": "CALC_004"}.items() if key != "mapping_status"},
            {**template, "case_id": "CALC_005", "mapping_status": " PROMOTED "},
        ])
        path.write_text(json.dumps(entries, ensure_ascii=False), encoding="utf-8")
        ledger = path.with_name("external_execution_history.jsonl")
        ledger.write_text(
            "".join(
                json.dumps(
                    _external_ledger_record(case_id, "计算器", "620C_W6830"),
                    ensure_ascii=False,
                ) + "\n"
                for case_id in ("CALC_001", "CALC_003", "CALC_004", "CALC_005")
            ),
            encoding="utf-8",
        )

        self.assertEqual(
            [row["case_id"] for row in self.cases.list(state_filter="solidified")["items"]],
            ["CALC_001"],
        )
        self.assertEqual(
            [row["case_id"] for row in self.cases.list(state_filter="explored_unsolidified")["items"]],
            ["CALC_003", "CALC_004", "CALC_005"],
        )
        self.assertEqual(
            [row["case_id"] for row in self.cases.list(state_filter="externally_explored")["items"]],
            ["CALC_001", "CALC_003", "CALC_004", "CALC_005"],
        )

        self.test_history.create(
            job={
                "sheet": "计算器",
                "case_id": "CALC_003",
                "case": {"priority": "P0"},
                "started_at": "2026-08-15T10:00:00+08:00",
                "finished_at": "2026-08-15T10:01:00+08:00",
                "return_code": 0,
            },
            result={"verdict": "FAIL", "reason": "FAIL"},
            stdout="",
            stderr="",
        )
        self.assertEqual(
            [row["case_id"] for row in self.cases.list(state_filter="explored_unsolidified")["items"]],
            ["CALC_003", "CALC_004", "CALC_005"],
        )

    def test_agent_test_history_is_isolated_by_project(self) -> None:
        hardware_case = self.cases.get("计算器", "CALC_001", project="6202_W5230")
        run_id = self.test_history.create(
            job={
                "project": "6202_W5230",
                "sheet": "计算器",
                "case_id": "CALC_001",
                "case": hardware_case,
                "started_at": "2026-08-13T10:00:00+08:00",
                "finished_at": "2026-08-13T10:00:05+08:00",
                "return_code": 0,
            },
            result={"verdict": "PASS", "reason": "真机符合预期"},
            stdout="done",
            stderr="",
        )

        self.assertEqual(self.test_history.list("计算器", "CALC_001"), [])
        hardware_history = self.test_history.list(
            "计算器", "CALC_001", project="6202_W5230"
        )
        self.assertEqual([item["id"] for item in hardware_history], [run_id])
        self.assertEqual(hardware_history[0]["execution_target"], "hardware")
        self.assertTrue(
            (
                self.paths.test_history
                / "6202_W5230"
                / "计算器"
                / "CALC_001"
                / run_id
                / "run.json"
            ).is_file()
        )

    def test_case_map_list_reuses_history_summary_and_updates_after_new_run(self) -> None:
        def record(verdict: str) -> str:
            return self.test_history.create(
                job={
                    "sheet": "计算器",
                    "case_id": "CALC_001",
                    "case": {"priority": "P0"},
                    "started_at": "2026-08-12T10:00:00+08:00",
                    "finished_at": "2026-08-12T10:01:00+08:00",
                    "return_code": 0,
                },
                result={"verdict": verdict, "reason": verdict},
                stdout="",
                stderr="",
            )

        first_id = record("FAIL")
        with patch.object(
            self.test_history,
            "_case_summary_from_dir",
            wraps=self.test_history._case_summary_from_dir,
        ) as summarize:
            first = self.cases.list()
            second = self.cases.list(state_filter="all")
            self.assertEqual(summarize.call_count, 1)
            self.assertEqual(first["items"][0]["latest_verdict"], "FAIL")
            self.assertEqual(second["items"][0]["history_count"], 1)
            with patch.object(
                self.cases,
                "_all",
                side_effect=AssertionError("recent 不应重建完整用例目录"),
            ):
                recent = self.cases.recent(limit=1)
            self.assertEqual(recent["items"][0]["case_id"], "CALC_001")
            self.assertEqual(recent["items"][0]["latest_verdict"], "FAIL")

            with patch.object(self.cases, "_all", wraps=self.cases._all) as load_all:
                overview = self.cases.overview(recent_limit=1, exception_limit=1)
            load_all.assert_called_once_with("620C_W6830")
            self.assertEqual(overview["recent_items"][0]["case_id"], "CALC_001")
            self.assertEqual(overview["recent_exceptions"][0]["latest_verdict"], "FAIL")

            second_id = record("PASS")
            updated = self.cases.list()
            self.assertEqual(summarize.call_count, 1)
            self.assertEqual(updated["items"][0]["latest_verdict"], "PASS")
            self.assertEqual(updated["items"][0]["history_count"], 2)
            self.assertEqual(updated["summary"]["solidified"], 1)
            self.assertEqual(self.cases.run_category(updated["items"][0]), "pass")

            self.assertTrue(self.test_history.delete("计算器", "CALC_001", second_id))
            after_delete = self.cases.list()
            self.assertEqual(summarize.call_count, 2)
            self.assertEqual(after_delete["items"][0]["latest_verdict"], "FAIL")
            self.assertEqual(after_delete["items"][0]["history_count"], 1)
            self.assertEqual(self.cases.run_category(after_delete["items"][0]), "fail")
        self.assertIsNotNone(self.test_history.get("计算器", "CALC_001", first_id))

    def test_promoted_case_with_unknown_historical_verdict_is_error(self) -> None:
        for verdict in (None, "MYSTERY"):
            with (
                self.subTest(verdict=verdict),
                patch.object(
                    self.test_history,
                    "summary_index",
                    return_value={
                        ("计算器", "CALC_001"): {
                            "latest": {
                                "verdict": verdict,
                                "timestamp": "2026-08-17T10:00:00+08:00",
                            },
                            "history_count": 1,
                        }
                    },
                ),
            ):
                row = self.cases.list()["items"][0]
                self.assertEqual(row["history_count"], 1)
                self.assertEqual(row["latest_verdict"], "ERROR")
                self.assertEqual(self.cases.run_category(row), "error")

    def test_batch_candidates_use_latest_history_and_exclude_pass(self) -> None:
        case_map_path = self.paths.case_map / "620C_simulator_case_map" / "计算器.json"
        entries = json.loads(case_map_path.read_text(encoding="utf-8"))
        template = dict(entries[0])
        for case_id in ("CALC_003", "CALC_004", "CALC_005", "CALC_006"):
            entries.append({**template, "case_id": case_id})
        case_map_path.write_text(json.dumps(entries, ensure_ascii=False), encoding="utf-8")

        def record(case_id: str, verdict: str) -> None:
            self.test_history.create(
                job={
                    "sheet": "计算器",
                    "case_id": case_id,
                    "case": {"priority": "P1"},
                    "started_at": "2026-08-11T10:00:00+08:00",
                    "finished_at": "2026-08-11T10:01:00+08:00",
                    "return_code": 0,
                },
                result={"verdict": verdict, "reason": verdict},
                stdout="",
                stderr="",
            )

        record("CALC_003", "FAIL")
        record("CALC_004", "CANNOT_VERIFY")
        record("CALC_005", "FAIL")
        record("CALC_005", "PASS")
        record("CALC_006", "ERROR")

        self.assertEqual(
            [row["case_id"] for row in self.cases.executable({"untested"})],
            ["CALC_001", "CALC_002"],
        )
        self.assertEqual(
            [row["case_id"] for row in self.cases.executable({"fail"})],
            ["CALC_003"],
        )
        self.assertEqual(
            [row["case_id"] for row in self.cases.executable({"cannot_verify"})],
            ["CALC_004"],
        )
        self.assertEqual(
            [row["case_id"] for row in self.cases.executable({"error"})],
            ["CALC_006"],
        )
        payload = self.cases.list()
        summary = payload["summary"]
        self.assertEqual(
            set(summary),
            {"all", "unexplored", "externally_explored", "explored_unsolidified", "solidified"},
        )
        self.assertEqual(
            payload["batch_summary"],
            {"untested": 2, "fail": 1, "cannot_verify": 1, "error": 1, "pass": 1},
        )

    def test_explicit_batch_selection_can_rerun_a_passed_case(self) -> None:
        case = self.cases.get("计算器", "CALC_001")
        self.assertIsNotNone(case)
        self.test_history.create(
            job={
                "sheet": "计算器",
                "case_id": "CALC_001",
                "case": case,
                "started_at": "2026-08-14T10:00:00+08:00",
                "finished_at": "2026-08-14T10:01:00+08:00",
                "return_code": 0,
            },
            result={"verdict": "PASS", "reason": "已通过"},
            stdout="",
            stderr="",
        )

        manager = CaseTestManager(self.paths, self.cases, self.test_history)
        with patch("frontend.server.threading.Thread.start") as start:
            job = manager.start_batch(
                case_refs=[{"sheet": "计算器", "case_id": "CALC_001"}],
            )

        queued_cases = manager._jobs[job["id"]]["cases"]
        self.assertEqual(job["total"], 1)
        self.assertEqual([item["case_id"] for item in queued_cases], ["CALC_001"])
        self.assertEqual(queued_cases[0]["latest_verdict"], "PASS")
        start.assert_called_once_with()

    def test_agent_test_history_saves_result_and_screenshot(self) -> None:
        screenshot = self.paths.evidence / "case.bmp"
        screenshot.write_bytes(b"BM-test")
        run_id = self.test_history.create(
            job={
                "sheet": "计算器",
                "case_id": "CALC_001",
                "started_at": "2026-08-11T10:00:00+08:00",
                "finished_at": "2026-08-11T10:00:05+08:00",
                "return_code": 0,
                "case": self.cases.get("计算器", "CALC_001"),
            },
            result={
                "verdict": "PASS",
                "reason": "GUI 树包含正确结果",
                "terminal_json": [{"type": "gui_tree_end", "status": "ok"}],
            },
            stdout="done",
            stderr="",
            screenshot=screenshot,
        )
        latest = self.test_history.latest("计算器", "CALC_001")
        self.assertEqual(latest["verdict"], "PASS")
        self.assertTrue(latest["has_screenshot"])
        record = self.test_history.get("计算器", "CALC_001", run_id)
        self.assertEqual(record["reason"], "GUI 树包含正确结果")
        self.assertEqual(record["terminal_json"][0]["status"], "ok")
        self.assertIn("screenshot_url", record)
        self.assertEqual(record["screenshot_urls"][0]["label"], "最终画面")

    def test_batch_records_many_groups_multiple_batches_in_one_scan(self) -> None:
        case = self.cases.get("计算器", "CALC_001")
        for batch_id, batch_token in (
            ("batch-one", "0001-CALC_001"),
            ("batch-two", "0001-CALC_001"),
        ):
            self.test_history.create(
                job={
                    "sheet": "计算器",
                    "case_id": "CALC_001",
                    "case": case,
                    "batch_id": batch_id,
                    "batch_token": batch_token,
                    "started_at": "2026-08-20T10:00:00+08:00",
                    "finished_at": "2026-08-20T10:00:01+08:00",
                    "return_code": 0,
                },
                result={"verdict": "PASS", "reason": "批次结果"},
                stdout="",
                stderr="",
            )

        records = self.test_history.batch_records_many({"batch-one", "batch-two"})

        self.assertEqual(set(records), {"batch-one", "batch-two"})
        self.assertEqual(set(records["batch-one"]), {"0001-CALC_001"})
        self.assertEqual(set(records["batch-two"]), {"0001-CALC_001"})
        self.assertEqual(
            self.test_history.batch_records("batch-one"),
            records["batch-one"],
        )
        with patch.object(
            self.test_history,
            "_case_summary_from_dir",
            side_effect=AssertionError("批次恢复扫描后不应再次读取历史目录"),
        ):
            summary_index = self.test_history.summary_index()
        self.assertEqual(summary_index[("计算器", "CALC_001")]["history_count"], 2)
        self.assertEqual(summary_index[("计算器", "CALC_001")]["latest"]["verdict"], "PASS")

    def test_agent_test_history_archives_multiple_checkpoint_screenshots(self) -> None:
        first = self.paths.evidence / "checkpoint-1.bmp"
        second = self.paths.evidence / "checkpoint-2.bmp"
        first.write_bytes(b"BM-one")
        second.write_bytes(b"BM-two")
        run_id = self.test_history.create(
            job={
                "sheet": "计算器",
                "case_id": "CALC_001",
                "started_at": "2026-08-11T10:00:00+08:00",
                "finished_at": "2026-08-11T10:00:05+08:00",
                "return_code": 0,
                "case": self.cases.get("计算器", "CALC_001"),
            },
            result={
                "verdict": "PASS",
                "reason": "两个状态均符合预期",
                "terminal_json": [],
                "planned_commands": {
                    "setup": [":ENTER_PAGE:CALCULATOR,0"],
                    "action": [":TP_CLICK:10,20,1"],
                    "collect": [":GUI_TREE:1"],
                },
                "command_trace": [
                    {
                        "index": 1,
                        "phase": "action",
                        "source": "case",
                        "kind": "device",
                        "wire": "srv_quick_cmd send TOP5STEP:TP_CLICK:10,20,1;",
                        "command": ":TP_CLICK:10,20,1",
                        "command_name": "TP_CLICK",
                        "status": "accepted",
                        "ok": True,
                    },
                    {
                        "index": 2,
                        "phase": "action",
                        "source": "runner",
                        "kind": "screenshot",
                        "wire": ":GUI_TREE:1",
                        "command": ":HOST_SCREENSHOT:GUI_TREE_CHECKPOINT",
                        "command_name": "HOST_SCREENSHOT",
                        "status": "completed",
                        "ok": True,
                        "checkpoint_index": 1,
                        "checkpoint_label": "状态一",
                    },
                ],
                "evidence_contract": {
                    "status": "COMPLETE",
                    "complete": True,
                    "required_screenshots": 2,
                    "captured_screenshots": 2,
                    "planned_action_count": 1,
                    "attempted_action_count": 1,
                    "business_action_count": 1,
                    "issues": [],
                },
                "screenshots": [
                    {"path": str(first), "label": "状态一", "phase": "action", "command": ":GUI_TREE:1", "trace_index": 2},
                    {"path": str(second), "label": "状态二", "phase": "action", "command": ":GUI_TREE:2", "trace_index": 4},
                ],
            },
            stdout="done",
            stderr="",
        )

        latest = self.test_history.latest("计算器", "CALC_001")
        self.assertEqual(latest["screenshot_count"], 2)
        record = self.test_history.get("计算器", "CALC_001", run_id)
        self.assertEqual(
            [item["label"] for item in record["screenshot_urls"]],
            ["状态一", "状态二"],
        )
        self.assertEqual(len(record["command_trace"]), 2)
        self.assertTrue(record["evidence_contract"]["complete"])
        self.assertEqual(record["screenshot_urls"][0]["trace_index"], 2)
        run_dir = self.test_history._run_dir("计算器", "CALC_001", run_id)
        self.assertEqual((run_dir / "screenshot-01.bmp").read_bytes(), b"BM-one")
        self.assertEqual((run_dir / "screenshot-02.bmp").read_bytes(), b"BM-two")

    def test_defect_detail_exposes_local_description_image(self) -> None:
        self._write_defect(
            "100",
            title="图片描述缺陷",
            description="[image]",
            attachments=[
                {
                    "name": "desc_image_0.png",
                    "kind": "image",
                    "summary": "本地缺陷原图",
                }
            ],
        )
        image_dir = self.paths.defect_images / "100"
        image_dir.mkdir(parents=True)
        image_bytes = b"\x89PNG\r\n\x1a\nlocal-image"
        (image_dir / "desc_image_0.png").write_bytes(image_bytes)

        detail = self.defects.get("100")
        self.assertEqual(
            detail["attachments"][0]["url"],
            "/api/defects/100/images/desc_image_0.png",
        )
        _, base = self._server()
        with urlopen(base + "/api/defects/100", timeout=3) as response:
            payload = json.loads(response.read().decode("utf-8"))
        self.assertEqual(payload["attachments"][0]["url"], "/api/defects/100/images/desc_image_0.png")
        with urlopen(base + payload["attachments"][0]["url"], timeout=3) as response:
            self.assertEqual(response.headers.get_content_type(), "image/png")
            self.assertEqual(response.read(), image_bytes)

    def test_defect_image_path_rejects_parent_segments(self) -> None:
        with self.assertRaisesRegex(ValueError, "图片名称格式不合法"):
            self.defects.defect_image_path("100", "..\\secret.png")

    def test_history_saves_agent_commands_actual_file_and_tristate(self) -> None:
        run_id = self._create_history()
        record = self.history.get("100", run_id)
        self.assertEqual(record["verdict"], "CANNOT_VERIFY")
        self.assertEqual(record["execution_mode"], "agent_generated")
        self.assertEqual(record["schema_version"], 2)
        self.assertFalse(record["legacy_record"])
        self.assertEqual(record["source_file"], "app/calculator.c")
        self.assertEqual(record["test_commands"], [":GUI_TREE:1", ":GET_CURRENT_WIN_ID"])
        self.assertEqual(record["verdict_reasons"], ["模拟器未返回完整数据"])
        self.assertEqual(record["llm_thinking_steps"], [])
        self.assertEqual(record["patch"]["after"], "new")
        self.assertIn("before", record["evidence"])
        self.assertEqual(self.defects.list()["items"][0]["repair_result"], "CANNOT_VERIFY")
        self.assertTrue(self.history.delete("100", run_id))
        self.assertIsNone(self.history.get("100", run_id))

    def test_history_reads_legacy_record_without_new_fields(self) -> None:
        self._write_history_record(
            "legacy",
            {"id": "legacy", "defect": "100", "verdict": "FAIL"},
        )
        record = self.history.get("100", "legacy")
        self.assertEqual(record["verdict"], "FAIL")
        self.assertTrue(record["legacy_record"])
        self.assertEqual(record["test_commands"], [])
        self.assertEqual(record["verdict_reasons"], [])
        self.assertEqual(record["llm_thinking_steps"], [])
        self.assertEqual(record["evidence"], {})

    def test_history_exposes_llm_thinking_steps_from_reproduction_trace(self) -> None:
        run_id = self.history.create(
            defect="100",
            job={"execution_mode": "agent_generated", "progress": {}},
            result={
                "verdict": "PASS",
                "reproduction_trace": {
                    "steps": [
                        {"step": 0, "decision": None},
                        {
                            "step": 1,
                            "decision": {
                                "action": "EXECUTE",
                                "command": ":ENTER_PAGE:VOICE_ASSISTANT_INTERACTION,0",
                                "reason": "先进入语音助手页面，确认初始界面。",
                            },
                        },
                        {
                            "step": 2,
                            "decision": {
                                "action": "EXECUTE",
                                "command": ":SIM_CONNECTION_SET:2,app,1",
                                "reason": "再建立连接，观察页面是否更新。",
                            },
                        },
                    ]
                },
            },
            stdout="",
            stderr="",
        )

        record = self.history.get("100", run_id)
        self.assertEqual(
            record["llm_thinking_steps"],
            ["先进入语音助手页面，确认初始界面。", "再建立连接，观察页面是否更新。"],
        )

    def test_history_saves_current_conforms_reason_and_commands_without_patch(self) -> None:
        run_id = self.history.create(
            defect="100",
            job={"execution_mode": "agent_generated", "progress": {}},
            result={
                "verdict": "PASS",
                "attempts": 0,
                "reproduction_outcome": "CURRENT_CONFORMS",
                "reproduction_reason": "当前界面符合预期，缺陷未复现。",
                "agent_test_commands": [":ENTER_PAGE:MESSAGE_LIST,0", ":ENTER_PAGE:SIDEBAR,0"],
                "patch": None,
                "test_output": None,
                "history": [],
                "patch_retained": False,
            },
            stdout="",
            stderr="",
        )

        record = self.history.get("100", run_id)
        self.assertEqual(record["schema_version"], 2)
        self.assertFalse(record["legacy_record"])
        self.assertEqual(record["verdict_reasons"], ["当前界面符合预期，缺陷未复现。"])
        self.assertEqual(
            record["test_commands"],
            [":ENTER_PAGE:MESSAGE_LIST,0", ":ENTER_PAGE:SIDEBAR,0"],
        )
        self.assertEqual(record["patch"], {})
        self.assertEqual(record["test_result"], {})

    def test_history_recovers_evidence_from_legacy_baseline_output(self) -> None:
        self._write_history_record(
            "legacy-current-conforms",
            {
                "id": "legacy-current-conforms",
                "defect": "100",
                "verdict": "PASS",
                "reproduction_outcome": "CURRENT_CONFORMS",
                "test_commands": [],
                "baseline_output": {
                    "reason": "旧格式中保存的判定理由",
                    "test_commands": [":ENTER_PAGE:SIDEBAR,0"],
                },
                "rounds": [],
            },
            test_result={},
        )

        record = self.history.get("100", "legacy-current-conforms")
        self.assertTrue(record["legacy_record"])
        self.assertEqual(record["verdict_reasons"], ["旧格式中保存的判定理由"])
        self.assertEqual(record["test_commands"], [":ENTER_PAGE:SIDEBAR,0"])

    def test_current_schema_missing_evidence_is_not_labeled_legacy(self) -> None:
        self._write_history_record(
            "current-incomplete",
            {
                "schema_version": 2,
                "id": "current-incomplete",
                "defect": "100",
                "verdict": "FAIL",
            },
        )

        record = self.history.get("100", "current-incomplete")
        self.assertFalse(record["legacy_record"])
        self.assertEqual(record["verdict_reasons"], [])
        self.assertEqual(record["test_commands"], [])
        self.assertEqual(record["llm_thinking_steps"], [])

    def test_history_recovers_commands_and_patch_from_previous_round(self) -> None:
        run_id = self.history.create(
            defect="100",
            job={"execution_mode": "agent_generated", "progress": {}},
            result={
                "verdict": "CANNOT_VERIFY",
                "patch": None,
                "test_output": None,
                "history": [
                    {
                        "patch": {"file_path": "app/calculator.c", "before": "old", "after": "new"},
                        "test_commands": [":GUI_TREE:1"],
                        "test_output": {"results": [{"verdict": "FAIL", "terminal_json": [{"ok": False}]}]},
                    }
                ],
            },
            stdout="",
            stderr="",
        )
        record = self.history.get("100", run_id)
        self.assertEqual(record["source_file"], "app/calculator.c")
        self.assertEqual(record["test_commands"], [":GUI_TREE:1"])
        self.assertEqual(record["patch"]["after"], "new")
        self.assertEqual(record["test_result"]["results"][0]["verdict"], "FAIL")

    def test_history_rejects_parent_directory_segments(self) -> None:
        with self.assertRaisesRegex(ValueError, "格式不合法"):
            self.history.get("..", "run")
        with self.assertRaisesRegex(ValueError, "格式不合法"):
            self.history.delete("100", "..")

    def test_defect_list_paginates_searches_and_clamps_last_page(self) -> None:
        for number in range(101, 125):
            self._write_defect(
                str(number),
                title=f"{'计算器' if number in {121, 122} else '活动记录'} 缺陷 {number}",
                description="按钮 显示异常" if number == 122 else "列表描述",
                status="处理中" if number in {121, 122} else "待处理",
            )

        first = self.defects.list()
        self.assertEqual(first["total"], 25)
        self.assertEqual(first["page_size"], 20)
        self.assertEqual(first["items"][0]["number"], "124")

        last = self.defects.list(page=999, page_size=10)
        self.assertEqual(last["page"], 3)
        self.assertEqual(len(last["items"]), 5)

        chinese = self.defects.list(query="计算器", page_size=10)
        self.assertEqual([row["number"] for row in chinese["items"]], ["122", "121", "100"])
        numbered = self.defects.list(query="122")
        self.assertEqual([row["number"] for row in numbered["items"]], ["122"])
        multiple = self.defects.list(query="  计算器   处理中 ")
        self.assertEqual([row["number"] for row in multiple["items"]], ["122", "121", "100"])
        empty = self.defects.list(query="完全不存在")
        self.assertEqual(empty["items"], [])
        self.assertEqual(empty["total_pages"], 0)

    def test_defect_list_filters_by_latest_repair_result(self) -> None:
        for number, verdict in (("101", "PASS"), ("102", "FAIL"), ("103", "CANNOT_VERIFY")):
            self._write_defect(number, title=f"筛选测试 {number}")
            self._create_history(verdict, defect=number)

        self.assertEqual([row["number"] for row in self.defects.list(result_filter="pass")["items"]], ["101"])
        self.assertEqual([row["number"] for row in self.defects.list(result_filter="fail")["items"]], ["102"])
        self.assertEqual([row["number"] for row in self.defects.list(result_filter="cannot_verify")["items"]], ["103"])
        self.assertEqual([row["number"] for row in self.defects.list(result_filter="pending")["items"]], ["100"])
        with self.assertRaisesRegex(ValueError, "result 参数不合法"):
            self.defects.list(result_filter="unknown")

    def test_http_api_page_routes_and_invalid_pagination(self) -> None:
        _, base = self._server()
        query = urlencode({"q": "计算器", "page": 1, "page_size": 20})
        with urlopen(base + "/api/defects?" + query, timeout=3) as response:
            payload = json.loads(response.read().decode("utf-8"))
            self.assertEqual(payload["items"][0]["number"], "100")
            self.assertEqual(payload["total"], 1)
        with urlopen(base + "/defect/100", timeout=3) as response:
            self.assertIn("正常".encode("utf-8"), response.read())
        for path in ("/api/defects?page=0", "/api/defects?page_size=101", "/api/defects?page=abc", "/api/defects?result=unknown"):
            with self.assertRaises(HTTPError) as raised:
                urlopen(base + path, timeout=3)
            self.assertEqual(raised.exception.code, 400)
        with self.assertRaises(HTTPError) as raised:
            urlopen(base + "/api/defects/missing", timeout=3)
        self.assertEqual(raised.exception.code, 404)

    def test_agent_test_api_lists_opens_and_starts_cases(self) -> None:
        application, base = self._server()
        with urlopen(base + "/api/tests?state=all&page=1&page_size=20", timeout=3) as response:
            payload = json.loads(response.read().decode("utf-8"))
        self.assertEqual([item["case_id"] for item in payload["items"]], ["CALC_001", "CALC_002"])
        self.assertEqual(payload["module_counts"], {"计算器": 2})
        with urlopen(base + "/api/tests?module=%E4%B8%8D%E5%AD%98%E5%9C%A8", timeout=3) as response:
            missing_module = json.loads(response.read().decode("utf-8"))
        self.assertEqual(missing_module["items"], [])
        self.assertEqual(missing_module["summary"]["all"], 0)
        with urlopen(base + "/api/tests/projects", timeout=3) as response:
            projects = json.loads(response.read().decode("utf-8"))
        self.assertEqual(len(projects["items"]), 3)
        with urlopen(base + "/api/tests/overview", timeout=3) as response:
            overview = json.loads(response.read().decode("utf-8"))
        self.assertEqual(overview["catalog_total"], 2)
        self.assertEqual(overview["recent_items"], [])
        with urlopen(base + "/api/tests/recent?limit=8", timeout=3) as response:
            recent = json.loads(response.read().decode("utf-8"))
        self.assertEqual(recent["items"], [])
        with urlopen(base + "/api/tests/%E8%AE%A1%E7%AE%97%E5%99%A8/CALC_001", timeout=3) as response:
            detail = json.loads(response.read().decode("utf-8"))
        self.assertEqual(detail["expected_text"], "显示正确")
        with urlopen(
            base + "/api/tests?project=6202_W5230&state=all&page=1&page_size=20",
            timeout=3,
        ) as response:
            hardware = json.loads(response.read().decode("utf-8"))
        self.assertEqual(hardware["execution_target"], "hardware")
        self.assertEqual(hardware["items"][0]["expected_text"], "真机显示正确")
        for page_path in ("/tests", "/test/%E8%AE%A1%E7%AE%97%E5%99%A8/CALC_001"):
            with urlopen(base + page_path, timeout=3) as response:
                self.assertIn("正常".encode("utf-8"), response.read())

        job = {"id": "test1", "sheet": "计算器", "case_id": "CALC_001", "status": "queued"}
        with patch.object(application, "start_case_test", return_value=job) as start:
            with self._post_json(base + "/api/tests/run", {"sheet": "计算器", "case_id": "CALC_001"}) as response:
                started = json.loads(response.read().decode("utf-8"))
        self.assertEqual(response.status, 202)
        self.assertEqual(started["id"], "test1")
        start.assert_called_once_with(
            sheet="计算器", case_id="CALC_001", project="620C_W6830"
        )

        hardware_job = {
            "id": "test2",
            "project": "6202_W5230",
            "execution_target": "hardware",
            "sheet": "计算器",
            "case_id": "CALC_001",
            "status": "queued",
        }
        with patch.object(application, "start_case_test", return_value=hardware_job) as start_hardware:
            with self._post_json(
                base + "/api/tests/run",
                {"project": "6202_W5230", "sheet": "计算器", "case_id": "CALC_001"},
            ) as response:
                started_hardware = json.loads(response.read().decode("utf-8"))
        self.assertEqual(started_hardware["execution_target"], "hardware")
        start_hardware.assert_called_once_with(
            sheet="计算器", case_id="CALC_001", project="6202_W5230"
        )

    def test_agent_test_api_starts_batch_and_serves_batch_page(self) -> None:
        application, base = self._server()
        job = {"id": "batch1", "type": "batch", "status": "queued", "total": 1}
        with patch.object(application, "start_batch_test", return_value=job) as start:
            with self._post_json(base + "/api/tests/run-batch", {"limit": 0}) as response:
                started = json.loads(response.read().decode("utf-8"))
        self.assertEqual(response.status, 202)
        self.assertEqual(started["id"], "batch1")
        start.assert_called_once_with(limit=0, project="620C_W6830")

        with patch.object(application, "start_batch_test", return_value=job) as start_selected:
            with self._post_json(base + "/api/tests/run-batch", {
                "cases": [{"sheet": "计算器", "case_id": "CALC_001"}]
            }) as response:
                selected = json.loads(response.read().decode("utf-8"))
        self.assertEqual(response.status, 202)
        self.assertEqual(selected["id"], "batch1")
        start_selected.assert_called_once_with(
            limit=0,
            case_refs=[{"sheet": "计算器", "case_id": "CALC_001"}],
            project="620C_W6830",
        )

        with patch.object(application, "start_batch_test", return_value=job) as start_categories:
            with self._post_json(base + "/api/tests/run-batch", {
                "categories": ["untested", "fail", "cannot_verify", "error"]
            }) as response:
                selected = json.loads(response.read().decode("utf-8"))
        self.assertEqual(response.status, 202)
        self.assertEqual(selected["id"], "batch1")
        start_categories.assert_called_once_with(
            limit=0,
            categories={"untested", "fail", "cannot_verify", "error"},
            project="620C_W6830",
        )

        with self.assertRaises(HTTPError) as raised:
            self._post_json(base + "/api/tests/run-batch", {"categories": ["pass"]})
        self.assertEqual(raised.exception.code, 400)
        with urlopen(base + "/test-batch/batch1", timeout=3) as response:
            self.assertIn("正常".encode("utf-8"), response.read())

        with self.assertRaises(HTTPError) as raised:
            self._post_json(base + "/api/tests/run-batch", {"limit": -1})
        self.assertEqual(raised.exception.code, 400)

    def test_case_test_manager_allows_one_hardware_and_one_simulator_job(self) -> None:
        manager = CaseTestManager(self.paths, self.cases, self.test_history)

        with patch("frontend.server.threading.Thread.start"):
            simulator = manager.start_batch(
                limit=1,
                project="6202_W5230_SIMULATOR",
            )
            hardware = manager.start_batch(
                limit=1,
                project="6202_W5230",
            )

            with self.assertRaisesRegex(RuntimeError, "模拟器资源已占用"):
                manager.start_batch(limit=1, project="620C_W6830")
            with self.assertRaisesRegex(RuntimeError, "真机资源已占用"):
                manager.start(
                    sheet="计算器",
                    case_id="CALC_001",
                    project="6202_W5230",
                )

        active = {
            job["execution_target"]: job
            for job in manager.active_jobs()
        }
        self.assertEqual(active["simulator"]["id"], simulator["id"])
        self.assertEqual(active["hardware"]["id"], hardware["id"])

    def test_active_test_api_returns_both_resource_slots(self) -> None:
        application, base = self._server()
        jobs = [
            {
                "id": "hardware-live",
                "project": "6202_W5230",
                "execution_target": "hardware",
                "status": "running",
            },
            {
                "id": "simulator-live",
                "project": "6202_W5230_SIMULATOR",
                "execution_target": "simulator",
                "status": "running",
            },
        ]
        with patch.object(
            application.test_jobs,
            "active_jobs",
            return_value=jobs,
        ) as active_jobs:
            with urlopen(base + "/api/tests/active", timeout=3) as response:
                payload = json.loads(response.read().decode("utf-8"))

        self.assertEqual(response.status, 200)
        self.assertEqual(payload["job"], jobs[0])
        self.assertEqual(payload["jobs"], jobs)
        active_jobs.assert_called_once_with()

    def test_case_test_manager_reuses_cli_and_archives_result(self) -> None:
        manager = CaseTestManager(self.paths, self.cases, self.test_history)
        job_id = "case-job"
        manager._jobs[job_id] = {
            "id": job_id,
            "sheet": "计算器",
            "case_id": "CALC_001",
            "case": self.cases.get("计算器", "CALC_001"),
            "status": "queued",
            "current_node": "load",
            "nodes": {node: "pending" for node in ("load", "execute", "judge", "record")},
            "verdict": "PENDING",
            "reason": "",
            "created_at": "2026-08-11T10:00:00+08:00",
            "started_at": None,
            "finished_at": None,
            "history_id": None,
            "error": None,
        }
        manager._active_job_ids["simulator"] = job_id

        class FakeProcess:
            returncode = 0

            def __init__(self, argv: list[str]):
                self.argv = argv

            def communicate(self) -> tuple[str, str]:
                result_file = Path(self.argv[self.argv.index("--result-file") + 1])
                screenshot = Path(self.argv[self.argv.index("--screenshot-path") + 1])
                result_file.write_text(
                    json.dumps({
                        "verdict": "PASS", "reason": "符合预期", "terminal_json": [{"status": "ok"}],
                        "setup_errors": [], "action_errors": [], "collect_errors": [],
                    }, ensure_ascii=False),
                    encoding="utf-8",
                )
                screenshot.write_bytes(b"BM-test")
                return "PASS", ""

        with patch("frontend.server.subprocess.Popen", side_effect=lambda argv, **_: FakeProcess(argv)):
            manager._run(job_id)

        snapshot = manager.get(job_id)
        self.assertEqual(snapshot["status"], "completed")
        self.assertEqual(snapshot["verdict"], "PASS")
        self.assertEqual(snapshot["nodes"], {"load": "pass", "execute": "pass", "judge": "pass", "record": "pass"})
        self.assertIsNotNone(snapshot["history_id"])

    def test_case_test_manager_uses_hardware_target_for_6202(self) -> None:
        manager = CaseTestManager(self.paths, self.cases, self.test_history)
        case = self.cases.get("计算器", "CALC_001", project="6202_W5230")
        manager._jobs["hardware-job"] = {"process": None}
        captured_argv: list[str] = []
        captured_env: dict[str, str] = {}

        class FakeProcess:
            returncode = 0

            def __init__(self, argv: list[str], env: dict[str, str]):
                captured_argv.extend(argv)
                captured_env.update(env)

            def communicate(self) -> tuple[str, str]:
                result_file = Path(captured_argv[captured_argv.index("--result-file") + 1])
                result_file.write_text(
                    json.dumps({"verdict": "PASS", "reason": "真机通过"}, ensure_ascii=False),
                    encoding="utf-8",
                )
                return "PASS", ""

        hardware_root = r"D:\Agent-loop-workspace\6202_W5230"
        profile_root = r"D:\Agent-loop\profiles"
        with patch.dict(os.environ, {
            "W30_SOURCE_ROOT": r"D:\Agent-loop-workspace\620C_W6830",
            "W30_AGENT_WORKSPACE_ROOT": r"D:\Agent-loop-workspace\620C_W6830",
            "W30_PROJECT": "620C_W6830",
            "W30_HARDWARE_SOURCE_ROOT": hardware_root,
            "W30_HARDWARE_WORKSPACE_ROOT": hardware_root,
            "W30_HARDWARE_PROJECT": "6202_W5230",
            "W30_HARDWARE_PROFILE_ROOT": profile_root,
        }, clear=False):
            with (
                patch("agent_loop_system.main._load_env") as load_env,
                patch(
                    "frontend.server.subprocess.Popen",
                    side_effect=lambda argv, **kwargs: FakeProcess(
                        argv, kwargs["env"]
                    ),
                ),
            ):
                result = manager._execute_case(
                    job_id="hardware-job",
                    case=case,
                    job_dir=self.paths.runtime_jobs / "hardware-job" / "single",
                )

        load_env.assert_called_once_with()
        self.assertEqual(captured_argv[captured_argv.index("--target") + 1], "hardware")
        self.assertNotIn("--preserve-test-session", captured_argv)
        self.assertNotIn("--skip-hardware-reset", captured_argv)
        self.assertNotIn("W30_SOURCE_ROOT", captured_env)
        self.assertNotIn("W30_AGENT_WORKSPACE_ROOT", captured_env)
        self.assertNotIn("W30_HARDWARE_SOURCE_ROOT", captured_env)
        self.assertNotIn("W30_HARDWARE_WORKSPACE_ROOT", captured_env)
        self.assertEqual(captured_env["W30_PROJECT"], "6202_W5230")
        self.assertEqual(captured_env["W30_HARDWARE_PROJECT"], "6202_W5230")
        self.assertEqual(captured_env["W30_HARDWARE_PROFILE_ROOT"], profile_root)
        self.assertEqual(result["verdict"], "PASS")

    def test_candidate_replay_job_passes_explicit_runner_flag(self) -> None:
        self._clear_case_candidate(project_dir="6202_case_map")
        context = self.cases.stage_agent_candidate(
            sheet="计算器",
            case_id="CALC_001",
            project="6202_W5230",
            source_history=self._agent_exploration_source(),
        )
        manager = CaseTestManager(self.paths, self.cases, self.test_history)
        case = self.cases.get("计算器", "CALC_001", project="6202_W5230")
        manager._jobs["candidate-cli"] = {
            "process": None,
            "candidate_replay": True,
        }
        captured_argv: list[str] = []

        class FakeProcess:
            returncode = 0

            def __init__(self, argv: list[str]):
                captured_argv.extend(argv)

            def communicate(self) -> tuple[str, str]:
                result_file = Path(captured_argv[captured_argv.index("--result-file") + 1])
                result_file.write_text(
                    json.dumps({"verdict": "PASS", "reason": "候选执行完成"}, ensure_ascii=False),
                    encoding="utf-8",
                )
                return "PASS", ""

        with patch(
            "frontend.server.subprocess.Popen",
            side_effect=lambda argv, **_: FakeProcess(argv),
        ):
            manager._execute_case(
                job_id="candidate-cli",
                case=case,
                job_dir=self.paths.runtime_jobs / "candidate-cli" / "single",
            )

        self.assertIn("--candidate-replay", captured_argv)
        self.cases.rollback_agent_candidate(context)

    def test_candidate_replay_start_stages_candidate_and_persists_recovery_state(self) -> None:
        case_map_path = self._clear_case_candidate(project_dir="6202_case_map")
        manager = CaseTestManager(self.paths, self.cases, self.test_history)

        with patch("frontend.server.threading.Thread.start"):
            job = manager.start(
                sheet="计算器",
                case_id="CALC_001",
                project="6202_W5230",
                candidate_replay=True,
                promotion_source=self._agent_exploration_source(),
            )

        self.assertTrue(job["promotion_flow"])
        self.assertTrue(job["candidate_replay"])
        self.assertEqual(job["promotion_status"], "pending")
        self.assertNotIn("promotion_context", job)
        state = json.loads(
            (self.paths.runtime_jobs / job["id"] / "promotion-state.json").read_text(
                encoding="utf-8"
            )
        )
        self.assertEqual(state["status"], "queued")
        staged = json.loads(case_map_path.read_text(encoding="utf-8"))["cases"][0]
        self.assertTrue(staged["actions"])
        rollback = self.cases.rollback_agent_candidate(
            manager._jobs[job["id"]]["promotion_context"]
        )
        self.assertEqual(rollback["status"], "rolled_back")

    def test_candidate_replay_job_auto_promotes_even_for_proven_product_fail(self) -> None:
        self._clear_case_candidate(project_dir="6202_case_map")
        ledger = self.paths.case_map / "6202_case_map" / "external_execution_history.jsonl"
        ledger.write_text("", encoding="utf-8")
        manager = CaseTestManager(self.paths, self.cases, self.test_history)
        context = self.cases.stage_agent_candidate(
            sheet="计算器",
            case_id="CALC_001",
            project="6202_W5230",
            source_history=self._agent_exploration_source(),
        )
        job_id = "auto-promotion"
        replay_result = self._candidate_replay_result(
            context,
            job_id=job_id,
            verdict="FAIL",
        )
        case = self.cases.get("计算器", "CALC_001", project="6202_W5230")
        manager._jobs[job_id] = {
            "id": job_id,
            "type": "single",
            "sheet": "计算器",
            "case_id": "CALC_001",
            "project": "6202_W5230",
            "project_label": "6202 W5230",
            "execution_target": "hardware",
            "execution_target_label": "真机",
            "case": case,
            "status": "queued",
            "current_node": "load",
            "nodes": {node: "pending" for node in ("load", "execute", "judge", "record")},
            "verdict": "PENDING",
            "reason": "",
            "created_at": "2026-08-21T12:00:00+08:00",
            "started_at": None,
            "finished_at": None,
            "history_id": None,
            "error": None,
            "candidate_replay": True,
            "promotion_flow": True,
            "promotion_status": "pending",
            "promotion_issues": [],
            "promotion_context": context,
        }
        manager._active_job_ids["hardware"] = job_id
        execution = {
            "result": replay_result,
            "stdout": "FAIL",
            "stderr": "",
            "return_code": 1,
            "verdict": "FAIL",
            "reason": replay_result["reason"],
            "execution_reason": "",
            "execute_failed": False,
            "started_at": "2026-08-21T12:00:00+08:00",
            "finished_at": "2026-08-21T12:01:00+08:00",
            "screenshot": self.paths.runtime_jobs / job_id / "single" / "screenshot.bmp",
        }

        with patch.object(manager, "_execute_case", return_value=execution):
            manager._run(job_id)

        snapshot = manager.get(job_id)
        self.assertEqual(snapshot["status"], "completed")
        self.assertEqual(snapshot["verdict"], "FAIL")
        self.assertEqual(snapshot["promotion_status"], "promoted")
        self.assertEqual(snapshot["promotion_issues"], [])
        promoted = self.cases.get("计算器", "CALC_001", project="6202_W5230")
        self.assertTrue(promoted["is_promoted"])
        self.assertFalse(promoted["external_explored"])
        self.assertEqual(ledger.read_text(encoding="utf-8"), "")

    def test_stale_candidate_is_rolled_back_when_frontend_restarts(self) -> None:
        case_map_path = self._clear_case_candidate(project_dir="6202_case_map")
        context = self.cases.stage_agent_candidate(
            sheet="计算器",
            case_id="CALC_001",
            project="6202_W5230",
            source_history=self._agent_exploration_source(),
        )
        job_id = "stale-promotion"
        context["job_id"] = job_id
        state_path = self.paths.runtime_jobs / job_id / "promotion-state.json"
        state_path.parent.mkdir(parents=True, exist_ok=True)
        state_path.write_text(json.dumps({
            "job_id": job_id,
            "status": "running",
            "context": context,
        }, ensure_ascii=False), encoding="utf-8")

        CaseTestManager(self.paths, self.cases, self.test_history)

        restored = json.loads(case_map_path.read_text(encoding="utf-8"))["cases"][0]
        self.assertNotIn("mapping_status", restored)
        self.assertEqual(restored["actions"], [])
        recovered_state = json.loads(state_path.read_text(encoding="utf-8"))
        self.assertEqual(recovered_state["status"], "rolled_back_on_restart")

    def test_case_test_manager_uses_isolated_6202_simulator_profile(self) -> None:
        manager = CaseTestManager(self.paths, self.cases, self.test_history)
        case = self.cases.get(
            "计算器", "CALC_001", project="6202_W5230_SIMULATOR"
        )
        manager._jobs["simulator-6202-job"] = {"process": None}
        captured_argv: list[str] = []
        captured_env: dict[str, str] = {}

        class FakeProcess:
            returncode = 0

            def __init__(self, argv: list[str], env: dict[str, str]):
                captured_argv.extend(argv)
                captured_env.update(env)

            def communicate(self) -> tuple[str, str]:
                result_file = Path(captured_argv[captured_argv.index("--result-file") + 1])
                result_file.write_text(
                    json.dumps({"verdict": "PASS", "reason": "截图通过"}, ensure_ascii=False),
                    encoding="utf-8",
                )
                return "PASS", ""

        simulator_root = (
            self.paths.root.parent
            / "workspaces"
            / "firmware"
            / "6202_W5230"
        ).resolve()
        with (
            patch.dict(
                os.environ,
                {
                    "AGENT_LOOP_ROOT": str(self.paths.root),
                    "W30_6202_SIMULATOR_SOURCE_ROOT": "../workspaces/firmware/6202_W5230",
                    "W30_6202_SIMULATOR_BUILD_DIRECTORY": "../workspaces/firmware/6202_W5230/core/gui/simulator/out/build/6202_W5230",
                    "W30_6202_SIMULATOR_ARTIFACT_PATH": "../workspaces/firmware/6202_W5230/core/gui/simulator/bin/main.exe",
                },
                clear=False,
            ),
            patch(
                "frontend.server.subprocess.Popen",
                side_effect=lambda argv, **kwargs: FakeProcess(argv, kwargs["env"]),
            ),
        ):
            result = manager._execute_case(
                job_id="simulator-6202-job",
                case=case,
                job_dir=self.paths.runtime_jobs / "simulator-6202-job" / "single",
            )

        self.assertEqual(captured_argv[captured_argv.index("--target") + 1], "simulator")
        self.assertEqual(
            captured_argv[captured_argv.index("--case-map-profile") + 1],
            "6202_W5230_SIMULATOR",
        )
        self.assertNotIn("--preserve-test-session", captured_argv)
        self.assertEqual(captured_env["W30_PROJECT"], "6202_W5230")
        self.assertEqual(
            captured_env["W30_SOURCE_ROOT"],
            str(simulator_root),
        )
        self.assertEqual(result["verdict"], "PASS")

    def test_collection_error_does_not_override_screenshot_verdict(self) -> None:
        manager = CaseTestManager(self.paths, self.cases, self.test_history)
        case = self.cases.get("计算器", "CALC_001", project="6202_W5230")
        manager._jobs["collect-warning"] = {"process": None}

        class FakeProcess:
            returncode = 0

            def __init__(self, argv: list[str]):
                self.argv = argv

            def communicate(self) -> tuple[str, str]:
                result_file = Path(self.argv[self.argv.index("--result-file") + 1])
                result_file.write_text(
                    json.dumps({
                        "verdict": "PASS",
                        "reason": "截图符合预期",
                        "setup_errors": [],
                        "action_errors": [],
                        "collect_errors": ["GUI_TREE 回包不完整"],
                    }, ensure_ascii=False),
                    encoding="utf-8",
                )
                return "PASS", ""

        with patch("frontend.server.subprocess.Popen", side_effect=lambda argv, **_: FakeProcess(argv)):
            result = manager._execute_case(
                job_id="collect-warning",
                case=case,
                job_dir=self.paths.runtime_jobs / "collect-warning" / "single",
            )

        self.assertEqual(result["verdict"], "PASS")
        self.assertTrue(result["execute_failed"])
        self.assertEqual(result["execution_reason"], "GUI_TREE 回包不完整")

    def test_incomplete_evidence_contract_marks_execution_failed(self) -> None:
        manager = CaseTestManager(self.paths, self.cases, self.test_history)
        case = self.cases.get("计算器", "CALC_001")
        manager._jobs["evidence-error"] = {"process": None}

        class FakeProcess:
            returncode = 1

            def __init__(self, argv: list[str]):
                self.argv = argv

            def communicate(self) -> tuple[str, str]:
                result_file = Path(self.argv[self.argv.index("--result-file") + 1])
                result_file.write_text(
                    json.dumps({
                        "schema_version": 2,
                        "verdict": "ERROR",
                        "reason": "业务动作未执行",
                        "evidence_contract": {
                            "complete": False,
                            "issues": [{
                                "code": "business_action_missing",
                                "message": "业务动作未执行",
                            }],
                        },
                        "setup_errors": [],
                        "action_errors": [],
                        "collect_errors": [],
                    }, ensure_ascii=False),
                    encoding="utf-8",
                )
                return "ERROR", ""

        with patch("frontend.server.subprocess.Popen", side_effect=lambda argv, **_: FakeProcess(argv)):
            result = manager._execute_case(
                job_id="evidence-error",
                case=case,
                job_dir=self.paths.runtime_jobs / "evidence-error" / "single",
            )

        self.assertEqual(result["verdict"], "ERROR")
        self.assertTrue(result["execute_failed"])
        self.assertEqual(result["execution_reason"], "业务动作未执行")

    def test_case_test_manager_runs_executable_batch_and_reports_progress(self) -> None:
        manager = CaseTestManager(self.paths, self.cases, self.test_history)
        job_id = "batch-job"
        case = self.cases.executable()[0]
        manager._jobs[job_id] = {
            "id": job_id,
            "type": "batch",
            "status": "queued",
            "created_at": "2026-08-11T10:00:00+08:00",
            "started_at": None,
            "finished_at": None,
            "total": 1,
            "completed": 0,
            "current_index": 0,
            "current_case": None,
            "current_node": "load",
            "verdict_counts": {"PASS": 0, "FAIL": 0, "ERROR": 0, "CANNOT_VERIFY": 0, "SKIP": 0},
            "recent_results": [],
            "cancel_requested": False,
            "error": None,
            "cases": [case],
        }
        manager._active_job_ids["simulator"] = job_id

        class FakeProcess:
            returncode = 0

            def __init__(self, argv: list[str]):
                self.argv = argv

            def communicate(self) -> tuple[str, str]:
                result_file = Path(self.argv[self.argv.index("--result-file") + 1])
                screenshot = Path(self.argv[self.argv.index("--screenshot-path") + 1])
                checkpoint = screenshot.with_name("screenshot-01.bmp")
                result_file.write_text(
                    json.dumps({
                        "verdict": "PASS",
                        "reason": "批次符合预期",
                        "terminal_json": [{"status": "ok"}],
                        "screenshots": [{"path": str(checkpoint), "label": "结果正确"}],
                        "setup_errors": [],
                        "action_errors": [],
                        "collect_errors": [],
                    }, ensure_ascii=False),
                    encoding="utf-8",
                )
                checkpoint.write_bytes(b"BM-checkpoint")
                return "PASS", ""

        with patch("frontend.server.subprocess.Popen", side_effect=lambda argv, **_: FakeProcess(argv)):
            manager._run_batch(job_id)

        snapshot = manager.get(job_id)
        self.assertEqual(snapshot["status"], "completed")
        self.assertEqual(snapshot["completed"], 1)
        self.assertEqual(snapshot["verdict_counts"]["PASS"], 1)
        self.assertEqual(snapshot["recent_results"][0]["case_id"], "CALC_001")
        self.assertIsNotNone(snapshot["recent_results"][0]["history_id"])
        self.assertEqual(len(snapshot["live_screenshots"]), 1)

    def test_case_test_manager_exposes_dynamic_exploration_screenshots(self) -> None:
        manager = CaseTestManager(self.paths, self.cases, self.test_history)
        job_id = "live-step-screenshots"
        token = "0001-CALC_001"
        current_dir = self.paths.runtime_jobs / job_id / token
        current_dir.mkdir(parents=True, exist_ok=True)
        (current_dir / "step_00.bmp").write_bytes(b"BM-step-0")
        (current_dir / "step_01.bmp").write_bytes(b"BM-step-1")
        (current_dir / "step_00_capture_original.bmp").write_bytes(b"BM-sidecar")
        manager._jobs[job_id] = {
            "id": job_id,
            "type": "batch",
            "status": "running",
            "total": 1,
            "completed": 0,
            "current_runtime_dir": str(current_dir),
            "current_case_data": {"verification_points": []},
            "current_case_token": token,
        }

        snapshot = manager.get(job_id)

        self.assertIsNotNone(snapshot)
        self.assertEqual(
            [item["label"] for item in snapshot["live_screenshots"]],
            ["探索步骤 1", "探索步骤 2"],
        )
        self.assertTrue(snapshot["live_screenshots"][0]["url"].endswith("/step_00.bmp"))
        self.assertEqual(
            manager.screenshot_path(job_id, token, "step_01.bmp"),
            current_dir / "step_01.bmp",
        )
        with self.assertRaisesRegex(ValueError, "截图文件名不合法"):
            manager.screenshot_path(job_id, token, "step_00_capture_original.bmp")
        with self.assertRaises(ValueError):
            manager.screenshot_path(job_id, token, "../step_00.bmp")

    def test_hardware_batch_stops_before_cases_when_case_reset_fails(self) -> None:
        manager = CaseTestManager(self.paths, self.cases, self.test_history)
        job_id = "hardware-preflight"
        case = self.cases.get("计算器", "CALC_001", project="6202_W5230")
        self.assertIsNotNone(case)
        manager._jobs[job_id] = {
            "id": job_id,
            "type": "batch",
            "project": "6202_W5230",
            "execution_target": "hardware",
            "status": "queued",
            "created_at": "2026-08-14T00:00:00+08:00",
            "started_at": None,
            "finished_at": None,
            "total": 1,
            "completed": 0,
            "current_index": 0,
            "current_case": None,
            "current_node": "load",
            "verdict_counts": {
                "PASS": 0,
                "FAIL": 0,
                "ERROR": 0,
                "CANNOT_VERIFY": 0,
                "SKIP": 0,
            },
            "recent_results": [],
            "cancel_requested": False,
            "error": None,
            "cases": [case],
        }
        manager._active_job_ids["hardware"] = job_id

        with (
            patch("agent_loop_system.main._load_env") as load_env,
            patch(
                "agent_loop_system.tools.hardware_runtime_profile.load_hardware_runtime_profile"
            ),
            patch(
                "agent_loop_system.tools.real_device.reset_hardware_case_state",
                side_effect=RuntimeError("GUI_STATE popup is CHARGING"),
            ) as reset,
            patch("frontend.server.subprocess.Popen") as popen,
        ):
            manager._run_batch(job_id)

        snapshot = manager.get(job_id)
        load_env.assert_called_once_with()
        reset.assert_called_once()
        popen.assert_not_called()
        self.assertEqual(snapshot["status"], "interrupted")
        self.assertEqual(snapshot["completed"], 0)
        self.assertIn("GUI_STATE popup is CHARGING", snapshot["error"])
        self.assertEqual(snapshot["hardware_reset"]["status"], "failed")
        self.assertNotIn("hardware", manager._active_job_ids)

    def test_hardware_batch_interrupts_on_infrastructure_failure_and_retries_case(self) -> None:
        manager = CaseTestManager(self.paths, self.cases, self.test_history)
        first = dict(self.cases.get("计算器", "CALC_001", project="6202_W5230"))
        second = {**first, "case_id": "CALC_002"}
        job_id = "hardware-interruption"
        manager._jobs[job_id] = {
            "id": job_id,
            "type": "batch",
            "project": "6202_W5230",
            "project_label": "6202 W5230",
            "execution_target": "hardware",
            "execution_target_label": "真机",
            "status": "queued",
            "created_at": "2026-08-14T00:00:00+08:00",
            "started_at": None,
            "finished_at": None,
            "total": 2,
            "completed": 0,
            "current_index": 0,
            "current_case": None,
            "current_node": "load",
            "verdict_counts": {
                "PASS": 0,
                "FAIL": 0,
                "ERROR": 0,
                "CANNOT_VERIFY": 0,
                "SKIP": 0,
            },
            "recent_results": [],
            "cancel_requested": False,
            "error": None,
            "case_attempts": {},
            "cases": [first, second],
        }
        manager._active_job_ids["hardware"] = job_id
        executed_cases: list[str] = []
        parent_reset_flags: list[bool] = []

        class BrokenProcess:
            returncode = 1

            def __init__(self, argv: list[str]):
                executed_cases.append(argv[argv.index("--case-id") + 1])
                parent_reset_flags.append("--skip-hardware-reset" in argv)

            def communicate(self) -> tuple[str, str]:
                return "", "HardwareSerialTimeoutError: no result for :GUI_PING:1"

        case_reset = SimpleNamespace(
            status=SimpleNamespace(active=True, lease_seconds=86000),
            reboot_status="accepted",
            gui_ping_attempts=1,
            bootstrap_event_seen=False,
            current_page="DIAL",
            popup=None,
        )
        with (
            patch("agent_loop_system.main._load_env"),
            patch("agent_loop_system.tools.hardware_runtime_profile.load_hardware_runtime_profile"),
            patch(
                "agent_loop_system.tools.real_device.reset_hardware_case_state",
                return_value=case_reset,
            ) as reset,
            patch(
                "frontend.server.subprocess.Popen",
                side_effect=lambda argv, **_: BrokenProcess(argv),
            ) as popen,
        ):
            manager._run_batch(job_id)

        interrupted = manager.get(job_id)
        self.assertEqual(reset.call_count, 1)
        self.assertEqual(popen.call_count, 1)
        self.assertEqual(executed_cases, ["CALC_001"])
        self.assertEqual(parent_reset_flags, [True])
        self.assertEqual(interrupted["status"], "interrupted")
        self.assertEqual(interrupted["completed"], 0)
        self.assertEqual(interrupted["current_index"], 1)
        self.assertTrue(interrupted["resume_available"])
        self.assertIn("CALC_001", interrupted["interruption_reason"])
        self.assertIn("GUI_PING", interrupted["error"])
        self.assertEqual(
            manager._jobs[job_id]["case_attempts"],
            {"0001-CALC_001": 1},
        )
        self.assertEqual(
            self.test_history.list("计算器", "CALC_001", project="6202_W5230"),
            [],
        )

        with patch("frontend.server.threading.Thread.start"):
            resumed = manager.resume_batch(job_id)
        self.assertEqual(resumed["status"], "queued")
        self.assertEqual(resumed["completed"], 0)
        self.assertIsNone(resumed["error"])

        class CompletedProcess:
            def __init__(self, argv: list[str]):
                self.argv = argv
                self.case_id = argv[argv.index("--case-id") + 1]
                self.returncode = 1 if self.case_id == "CALC_001" else 0
                executed_cases.append(self.case_id)
                parent_reset_flags.append("--skip-hardware-reset" in argv)

            def communicate(self) -> tuple[str, str]:
                result_file = Path(self.argv[self.argv.index("--result-file") + 1])
                screenshot = Path(self.argv[self.argv.index("--screenshot-path") + 1])
                checkpoint = screenshot.with_name("screenshot-01.bmp")
                verdict = "FAIL" if self.case_id == "CALC_001" else "PASS"
                result_file.write_text(json.dumps({
                    "verdict": verdict,
                    "reason": "产品结果已形成",
                    "screenshots": [{"path": str(checkpoint), "label": "结果"}],
                    "setup_errors": [],
                    "action_errors": [],
                    "collect_errors": [],
                }, ensure_ascii=False), encoding="utf-8")
                checkpoint.write_bytes(b"BM-completed")
                return verdict, ""

        with (
            patch("agent_loop_system.main._load_env"),
            patch("agent_loop_system.tools.hardware_runtime_profile.load_hardware_runtime_profile"),
            patch(
                "agent_loop_system.tools.real_device.reset_hardware_case_state",
                return_value=case_reset,
            ) as reset,
            patch(
                "frontend.server.subprocess.Popen",
                side_effect=lambda argv, **_: CompletedProcess(argv),
            ),
        ):
            manager._run_batch(job_id)

        completed = manager.get(job_id)
        self.assertEqual(reset.call_count, 2)
        self.assertEqual(executed_cases, ["CALC_001", "CALC_001", "CALC_002"])
        self.assertEqual(parent_reset_flags, [True, True, True])
        self.assertEqual(completed["status"], "completed")
        self.assertEqual(completed["completed"], 2)
        self.assertEqual(completed["verdict_counts"]["FAIL"], 1)
        self.assertEqual(completed["verdict_counts"]["PASS"], 1)
        self.assertEqual(
            manager._jobs[job_id]["case_attempts"],
            {"0001-CALC_001": 2, "0002-CALC_002": 1},
        )
        self.assertTrue(
            (self.paths.runtime_jobs / job_id / "0001-CALC_001").is_dir()
        )

    def test_batch_restart_scans_history_once_for_all_saved_batches(self) -> None:
        manager = CaseTestManager(self.paths, self.cases, self.test_history)
        case = dict(self.cases.executable()[0])
        job_ids = {"saved-batch-one", "saved-batch-two"}
        for job_id in job_ids:
            manager._jobs[job_id] = {
                "id": job_id,
                "type": "batch",
                "status": "failed",
                "created_at": "2026-08-20T00:00:00+08:00",
                "started_at": "2026-08-20T00:00:01+08:00",
                "finished_at": "2026-08-20T00:00:02+08:00",
                "total": 1,
                "completed": 0,
                "current_index": 0,
                "current_case": None,
                "current_node": None,
                "verdict_counts": {
                    "PASS": 0,
                    "FAIL": 0,
                    "ERROR": 0,
                    "CANNOT_VERIFY": 0,
                },
                "recent_results": [],
                "cancel_requested": False,
                "error": "测试中断",
                "case_attempts": {},
                "cases": [case],
            }
            with manager._lock:
                manager._persist_batch_locked(manager._jobs[job_id])

        with (
            patch.object(
                self.test_history,
                "batch_records_many",
                wraps=self.test_history.batch_records_many,
            ) as scan_many,
            patch.object(
                self.test_history,
                "batch_records",
                side_effect=AssertionError("启动恢复不应逐批扫描历史"),
            ),
        ):
            restarted = CaseTestManager(self.paths, self.cases, self.test_history)

        scan_many.assert_called_once_with(job_ids)
        self.assertEqual(set(restarted._jobs), job_ids)

    def test_batch_restart_recovers_same_id_and_runs_only_remaining_case(self) -> None:
        manager = CaseTestManager(self.paths, self.cases, self.test_history)
        first = dict(self.cases.executable()[0])
        second = {**first, "case_id": "CALC_002"}
        job_id = "restart-batch"
        manager._jobs[job_id] = {
            "id": job_id,
            "type": "batch",
            "status": "running",
            "created_at": "2026-08-12T00:00:00+08:00",
            "started_at": "2026-08-12T00:00:01+08:00",
            "finished_at": None,
            "total": 2,
            "completed": 1,
            "current_index": 2,
            "current_case": {"sheet": "计算器", "case_id": "CALC_002"},
            "current_node": "execute",
            "verdict_counts": {"PASS": 1, "FAIL": 0, "ERROR": 0, "CANNOT_VERIFY": 0, "SKIP": 0},
            "recent_results": [],
            "cancel_requested": False,
            "error": None,
            "cases": [first, second],
        }
        with manager._lock:
            manager._persist_batch_locked(manager._jobs[job_id])

        restarted = CaseTestManager(self.paths, self.cases, self.test_history)
        interrupted = restarted.get(job_id)
        self.assertEqual(interrupted["status"], "interrupted")
        self.assertTrue(interrupted["resume_available"])
        self.assertEqual(interrupted["completed"], 1)

        with patch("frontend.server.threading.Thread.start") as start:
            resumed = restarted.resume_batch(job_id)
            repeated = restarted.resume_batch(job_id)
        self.assertEqual(resumed["id"], job_id)
        self.assertEqual(repeated["id"], job_id)
        self.assertEqual(start.call_count, 1)

        executed_cases: list[str] = []

        class FakeProcess:
            returncode = 0

            def __init__(self, argv: list[str]):
                self.argv = argv
                executed_cases.append(argv[argv.index("--case-id") + 1])

            def communicate(self) -> tuple[str, str]:
                result_file = Path(self.argv[self.argv.index("--result-file") + 1])
                screenshot = Path(self.argv[self.argv.index("--screenshot-path") + 1])
                checkpoint = screenshot.with_name("screenshot-01.bmp")
                result_file.write_text(json.dumps({
                    "verdict": "PASS", "reason": "恢复后通过", "screenshots": [{"path": str(checkpoint)}],
                    "setup_errors": [], "action_errors": [], "collect_errors": [],
                }, ensure_ascii=False), encoding="utf-8")
                checkpoint.write_bytes(b"BM-resumed")
                return "PASS", ""

        with patch("frontend.server.subprocess.Popen", side_effect=lambda argv, **_: FakeProcess(argv)):
            restarted._run_batch(job_id)

        completed = restarted.get(job_id)
        self.assertEqual(executed_cases, ["CALC_002"])
        self.assertEqual(completed["status"], "completed")
        self.assertEqual(completed["completed"], 2)
        self.assertFalse(completed["resume_available"])

        loaded_again = CaseTestManager(self.paths, self.cases, self.test_history)
        self.assertEqual(loaded_again.get(job_id)["status"], "completed")
        with self.assertRaisesRegex(ValueError, "全部完成"):
            loaded_again.resume_batch(job_id)

    def test_cancelled_batch_is_resumable_without_resetting_progress(self) -> None:
        manager = CaseTestManager(self.paths, self.cases, self.test_history)
        case = dict(self.cases.executable()[0])
        job_id = "paused-batch"
        manager._jobs[job_id] = {
            "id": job_id, "type": "batch", "status": "cancelled",
            "created_at": "2026-08-12T00:00:00+08:00", "started_at": "2026-08-12T00:00:01+08:00",
            "finished_at": "2026-08-12T00:01:00+08:00", "total": 2, "completed": 1,
            "current_index": 1, "current_case": None, "current_node": None,
            "verdict_counts": {"PASS": 1, "FAIL": 0, "ERROR": 0, "CANNOT_VERIFY": 0, "SKIP": 0},
            "recent_results": [], "cancel_requested": True, "error": None,
            "interruption_reason": "用户请求在当前用例结束后暂停",
            "cases": [case, {**case, "case_id": "CALC_002"}],
        }
        with patch("frontend.server.threading.Thread.start") as start:
            snapshot = manager.resume_batch(job_id)
        self.assertEqual(snapshot["status"], "queued")
        self.assertEqual(snapshot["completed"], 1)
        self.assertEqual(snapshot["current_index"], 1)
        self.assertFalse(snapshot["cancel_requested"])
        start.assert_called_once_with()

    def test_resume_batch_api_uses_existing_batch(self) -> None:
        application, base = self._server()
        job = {"id": "batch1", "type": "batch", "status": "queued", "completed": 3, "total": 5}
        with patch.object(application, "resume_batch_test", return_value=job) as resume:
            with self._post_json(base + "/api/tests/jobs/batch1/resume", {}) as response:
                payload = json.loads(response.read().decode("utf-8"))
        self.assertEqual(response.status, 202)
        self.assertEqual(payload["id"], "batch1")
        resume.assert_called_once_with("batch1")

    def test_run_api_accepts_only_defect_and_returns_conflict(self) -> None:
        application, base = self._server()
        job = {"id": "job1", "defect": "100", "status": "queued", "verdict": "PENDING"}
        with patch.object(application.jobs, "start", return_value=job) as start:
            with self._post_json(base + "/api/run", {"defect": "100"}) as response:
                payload = json.loads(response.read().decode("utf-8"))
            self.assertEqual(response.status, 202)
            self.assertEqual(payload["id"], "job1")
            start.assert_called_once_with(defect="100")

        for payload in ({}, {"defect": ""}, {"defect": 100}):
            with self.assertRaises(HTTPError) as raised:
                self._post_json(base + "/api/run", payload)
            self.assertEqual(raised.exception.code, 400)

        with patch.object(application.jobs, "start", side_effect=RuntimeError("已有修复任务正在运行")):
            with self.assertRaises(HTTPError) as raised:
                self._post_json(base + "/api/run", {"defect": "100"})
            self.assertEqual(raised.exception.code, 409)

    def test_active_run_api_returns_current_job(self) -> None:
        application, base = self._server()
        job = {"id": "job-live", "defect": "100", "status": "running"}
        with patch.object(application.jobs, "active", return_value=job) as active:
            with urlopen(base + "/api/run/active", timeout=3) as response:
                payload = json.loads(response.read().decode("utf-8"))
        self.assertEqual(response.status, 200)
        self.assertEqual(payload["job"], job)
        active.assert_called_once_with()

    def test_active_job_snapshot_only_returns_live_job(self) -> None:
        manager = JobManager(self.paths, self.defects, self.history)
        job_id = "job-live"
        manager._active_job_id = job_id
        manager._jobs[job_id] = {
            "id": job_id,
            "defect": "100",
            "status": "running",
            "nodes": {node: "pending" for node in ("validate", "interactive_reproduce", "agent", "apply", "build", "test", "record")},
            "verdict": "PENDING",
        }
        self.assertEqual(manager.active()["id"], job_id)
        manager._jobs[job_id]["status"] = "finalizing"
        self.assertEqual(manager.active()["status"], "finalizing")
        manager._jobs[job_id]["status"] = "completed"
        self.assertIsNone(manager.active())

    def test_job_stays_finalizing_until_history_is_saved(self) -> None:
        manager = JobManager(self.paths, self.defects, self.history)
        job_id = "job-finalize"
        manager._jobs[job_id] = {
            "id": job_id,
            "defect": "100",
            "status": "queued",
            "current_node": None,
            "nodes": {
                node: "pending"
                for node in ("validate", "interactive_reproduce", "agent", "apply", "build", "test", "record")
            },
            "verdict": "PENDING",
            "created_at": "2026-08-08T15:00:00+08:00",
            "started_at": None,
            "finished_at": None,
            "history_id": None,
            "error": None,
        }
        manager._active_job_id = job_id
        history_started = threading.Event()
        allow_history_finish = threading.Event()

        class FakeProcess:
            returncode = 0

            def __init__(self, argv: list[str]):
                self.argv = argv

            def communicate(self) -> tuple[str, str]:
                result_file = Path(self.argv[self.argv.index("--result-file") + 1])
                result_file.write_text(
                    json.dumps({"verdict": "FAIL", "attempts": 1, "error": None}),
                    encoding="utf-8",
                )
                return "done", ""

        def save_history(**_: object) -> str:
            history_started.set()
            self.assertTrue(allow_history_finish.wait(timeout=3))
            return "history-new"

        with (
            patch("frontend.server.subprocess.Popen", side_effect=lambda argv, **_: FakeProcess(argv)),
            patch.object(manager.history, "create", side_effect=save_history),
        ):
            worker = threading.Thread(target=manager._run, args=(job_id,), daemon=True)
            worker.start()
            self.assertTrue(history_started.wait(timeout=3))
            snapshot = manager.get(job_id)
            self.assertEqual(snapshot["status"], "finalizing")
            self.assertIsNone(snapshot["history_id"])
            allow_history_finish.set()
            worker.join(timeout=3)

        snapshot = manager.get(job_id)
        self.assertEqual(snapshot["status"], "completed")
        self.assertEqual(snapshot["history_id"], "history-new")

    def test_finalizing_job_still_blocks_a_second_repair(self) -> None:
        manager = JobManager(self.paths, self.defects, self.history)
        manager._active_job_id = "active"
        manager._jobs["active"] = {"id": "active", "status": "finalizing"}
        with self.assertRaisesRegex(RuntimeError, "已有修复任务"):
            manager.start(defect="100")

    def test_import_api_accepts_options_reports_progress_and_returns_conflict(self) -> None:
        application, base = self._server()
        job = {
            "id": "import1", "type": "import", "status": "running",
            "include_completed": True, "limit": 50, "force": True,
        }
        with patch("frontend.server.RequestHandler._start_import", return_value=job) as start:
            with self._post_json(
                base + "/api/defects/import",
                {"include_completed": True, "limit": 50, "force": True},
            ) as response:
                payload = json.loads(response.read().decode("utf-8"))
            self.assertEqual(response.status, 202)
            self.assertEqual(payload["id"], "import1")
            start.assert_called_once_with(include_completed=True, limit=50, force=True)

        for payload in (
            {"include_completed": "true", "limit": 0, "force": True},
            {"include_completed": False, "limit": -1, "force": True},
            {"include_completed": False, "limit": 0, "force": 1},
        ):
            with self.assertRaises(HTTPError) as raised:
                self._post_json(base + "/api/defects/import", payload)
            self.assertEqual(raised.exception.code, 400)

        with patch(
            "frontend.server.RequestHandler._start_import",
            side_effect=RuntimeError("已有导入任务 import1 正在运行"),
        ):
            with self.assertRaises(HTTPError) as raised:
                self._post_json(
                    base + "/api/defects/import",
                    {"include_completed": False, "limit": 0, "force": True},
                )
            self.assertEqual(raised.exception.code, 409)

        log_file = self.paths.runtime_jobs / "import1" / "import.log"
        log_file.parent.mkdir(parents=True, exist_ok=True)
        log_file.write_text("[import] 处理缺陷 100\n[import] 完成\n", encoding="utf-8")
        application.import_jobs["import1"] = {
            **job, "log_file": str(log_file), "finished_at": None, "error": None,
        }
        with urlopen(base + "/api/defects/import/import1", timeout=3) as response:
            progress = json.loads(response.read().decode("utf-8"))
        self.assertIn("处理缺陷 100", progress["stdout_tail"])

    def test_import_runner_respects_force_include_completed_and_limit(self) -> None:
        application = WebApplication(self.paths)
        handler = object.__new__(RequestHandler)
        handler.app = application
        log_file = self.paths.runtime_jobs / "runner" / "import.log"
        log_file.parent.mkdir(parents=True, exist_ok=True)

        application.import_jobs["runner"] = {"id": "runner", "status": "running"}
        with patch("frontend.server.subprocess.run", return_value=SimpleNamespace(returncode=0)) as run:
            handler._run_import("runner", True, 50, False, str(log_file))
        argv = run.call_args.args[0]
        self.assertIn("--include-completed", argv)
        self.assertEqual(argv[argv.index("--limit") + 1], "50")
        self.assertNotIn("--force", argv)

        application.import_jobs["runner2"] = {"id": "runner2", "status": "running"}
        with patch("frontend.server.subprocess.run", return_value=SimpleNamespace(returncode=0)) as run:
            handler._run_import("runner2", False, 0, True, str(log_file))
        argv = run.call_args.args[0]
        self.assertIn("--force", argv)
        self.assertNotIn("--include-completed", argv)
        self.assertNotIn("--limit", argv)



    def test_case_crud_and_export_endpoints(self) -> None:
        _, base = self._server()
        
        # 1. Create Case
        create_payload = {
            "project": "620C_W6830",
            "case": {
                "case_id": "CALC_NEW_01",
                "sheet": "计算器",
                "priority": "P1",
                "precondition_text": "已打开计算器",
                "steps_text": "点击按键1",
                "expected_text": "屏幕显示1",
                "note": "自动化新建用例",
            }
        }
        with self._post_json(base + "/api/cases/create", create_payload) as resp:
            data = json.loads(resp.read().decode("utf-8"))
            self.assertEqual(resp.status, 200)
            self.assertEqual(data["case_id"], "CALC_NEW_01")
            
        # Verify duplicate creation error
        with self.assertRaises(HTTPError) as err:
            self._post_json(base + "/api/cases/create", create_payload)
        self.assertEqual(err.exception.code, 400)
        
        # 2. Update Case
        update_payload = {
            "project": "620C_W6830",
            "orig_case_id": "CALC_NEW_01",
            "case": {
                "case_id": "CALC_NEW_01",
                "sheet": "计算器",
                "priority": "P0",
                "precondition_text": "已打开计算器并重置",
                "steps_text": "点击按键1与按键2",
                "expected_text": "屏幕显示12",
                "note": "更新备注",
            }
        }
        with self._post_json(base + "/api/cases/update", update_payload) as resp:
            data = json.loads(resp.read().decode("utf-8"))
            self.assertEqual(resp.status, 200)
            self.assertEqual(data["status"], "ok")
            
        # 3. Export Cases XLSX
        with urlopen(base + "/api/cases/export?project=620C_W6830", timeout=3) as resp:
            self.assertEqual(resp.status, 200)
            self.assertIn("application/vnd.openxmlformats-officedocument", resp.headers.get("Content-Type"))
            xlsx_bytes = resp.read()
            self.assertTrue(len(xlsx_bytes) > 1000)
            import openpyxl
            import io
            wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
            self.assertIn("自动化测试用例_v1", wb.sheetnames)
            ws = wb["自动化测试用例_v1"]
            self.assertEqual(ws.freeze_panes, "A2")
            self.assertEqual(ws.auto_filter.ref, f"A1:I{ws.max_row}")
            self.assertFalse(ws.sheet_view.showGridLines)
            self.assertEqual(ws.page_setup.orientation, "landscape")
            self.assertEqual(ws.column_dimensions["E"].width, 46)
            self.assertEqual(ws.row_dimensions[1].height, 28)

            header = ws["A1"]
            self.assertEqual(header.font.name, "宋体")
            self.assertEqual(header.font.sz, 11)
            self.assertTrue(header.font.bold)
            self.assertEqual(header.fill.fgColor.rgb[-6:], "1F4E78")
            self.assertEqual(header.border.left.style, "thin")

            case_row = next(row for row in ws.iter_rows(min_row=2) if row[1].value == "CALC_NEW_01")
            self.assertEqual(case_row[0].font.name, "宋体")
            self.assertEqual(case_row[1].font.name, "Times New Roman")
            self.assertEqual(case_row[1].font.sz, 10)
            self.assertTrue(case_row[4].alignment.wrap_text)
            self.assertEqual(case_row[4].alignment.vertical, "top")
            self.assertEqual(case_row[4].border.bottom.style, "thin")
            self.assertGreaterEqual(ws.row_dimensions[case_row[0].row].height, 22)

    def test_excel_import_preview_and_confirm(self) -> None:
        _, base = self._server()
        import openpyxl
        import io
        import base64
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "自动化测试用例_v1"
        ws.append(["模块/Sheet", "用例编号", "优先级", "前置条件", "测试步骤", "预期结果", "不可自动化", "固化状态", "备注"])
        ws.append(["控制中心", "CTRL_001", "P0", "在主表盘下滑", "查看控制中心", "显示WiFi和蓝牙开关", "否", "", "测试导入"])
        ws.append(["计算器", "CALC_001", "P0", "原前置", "原步骤", "原预期", "否", "", "已有用例"])
        
        buf = io.BytesIO()
        wb.save(buf)
        b64 = base64.b64encode(buf.getvalue()).decode("utf-8")
        
        # 1. Preview
        preview_req = {
            "project": "620C_W6830",
            "file_name": "test_import.xlsx",
            "file_base64": b64,
        }
        with self._post_json(base + "/api/excel/preview", preview_req) as resp:
            data = json.loads(resp.read().decode("utf-8"))
            self.assertEqual(resp.status, 200)
            self.assertEqual(data["total_parsed"], 2)
            self.assertEqual(data["new_count"], 1)
            self.assertEqual(data["existing_count"], 1)
            self.assertIn("控制中心", data["modules"])
            
        # 2. Confirm without overwrite
        confirm_req = {
            "project": "620C_W6830",
            "file_name": "test_import.xlsx",
            "file_base64": b64,
            "overwrite_existing": False,
        }
        with self._post_json(base + "/api/excel/confirm", confirm_req) as resp:
            data = json.loads(resp.read().decode("utf-8"))
            self.assertEqual(resp.status, 200)
            self.assertEqual(data["imported_count"], 1)
            self.assertEqual(data["skipped_count"], 1)
            self.assertIn("控制中心", data["modules_updated"])

    def test_migration_and_audit_endpoints(self) -> None:
        application, base = self._server()
        
        # 1. Migration Candidates
        with urlopen(base + "/api/cases/migration-candidates?source=6202_W5230_SIMULATOR&target=6202_W5230", timeout=3) as resp:
            data = json.loads(resp.read().decode("utf-8"))
            self.assertEqual(resp.status, 200)
            self.assertIn("candidates", data)
            
        # 2. Audit and Promote
        audit_payload = {
            "case_id": "CALC_001",
            "sheet": "计算器",
            "project": "620C_W6830",
        }
        case_map_path = self.paths.case_map / "620C_simulator_case_map" / "计算器.json"
        case_map = json.loads(case_map_path.read_text(encoding="utf-8"))
        case_map[0]["mapping_status"] = ""
        case_map_path.write_text(json.dumps(case_map, ensure_ascii=False), encoding="utf-8")
        # Initially no history run -> returns 400 with audit issues
        with self.assertRaises(HTTPError) as err:
            self._post_json(base + "/api/cases/audit-and-promote", audit_payload)
        self.assertEqual(err.exception.code, 400)

        # A complete Agent exploration starts an explicit candidate replay; the
        # endpoint itself no longer writes PROMOTED.
        application.test_history._create(
            job={
                "sheet": "计算器",
                "case_id": "CALC_001",
                "project": "620C_W6830",
                "finished_at": "2026-08-20T10:00:00+08:00",
            },
            result={
                "verdict": "PASS",
                "reason": "自主探索完成",
                "execution_mode": "agent_exploration",
                "evidence_contract": {"status": "COMPLETE", "complete": True},
            },
            stdout="PASS",
            stderr="",
        )
        candidate_job = {
            "id": "candidate-job",
            "status": "queued",
            "promotion_flow": True,
        }
        with patch.object(
            application,
            "start_candidate_replay",
            return_value=candidate_job,
        ) as start_candidate:
            with self._post_json(base + "/api/cases/audit-and-promote", audit_payload) as resp:
                data = json.loads(resp.read().decode("utf-8"))
        self.assertEqual(resp.status, 202)
        self.assertEqual(data["status"], "candidate_replay_started")
        self.assertEqual(data["job"], candidate_job)
        call = start_candidate.call_args.kwargs
        self.assertEqual(call["sheet"], "计算器")
        self.assertEqual(call["case_id"], "CALC_001")
        self.assertEqual(call["project"], "620C_W6830")
        self.assertEqual(call["source_history"]["execution_mode"], "agent_exploration")
        case_map = json.loads(case_map_path.read_text(encoding="utf-8"))
        self.assertEqual(case_map[0]["mapping_status"], "")

        # A completed flow is idempotent and still does not append the ledger.
        case_map[0]["mapping_status"] = "PROMOTED"
        case_map_path.write_text(json.dumps(case_map, ensure_ascii=False), encoding="utf-8")
        with self._post_json(base + "/api/cases/audit-and-promote", audit_payload) as resp:
            data = json.loads(resp.read().decode("utf-8"))
            self.assertEqual(resp.status, 200)
            self.assertTrue(data["audit"]["passed"])
            self.assertEqual(data["status"], "already_promoted")

        # Repeated requests must not duplicate the external exploration record.
        with self._post_json(base + "/api/cases/audit-and-promote", audit_payload) as resp:
            self.assertEqual(resp.status, 200)
        ledger = self.paths.case_map / "620C_simulator_case_map" / "external_execution_history.jsonl"
        ledger_case_ids = [
            json.loads(line)["case_id"]
            for line in ledger.read_text(encoding="utf-8").splitlines()
            if line.strip()
        ]
        self.assertEqual(ledger_case_ids.count("CALC_001"), 1)
        with urlopen(base + "/api/tests?project=620C_W6830", timeout=3) as resp:
            self.assertEqual(resp.status, 200)

    def test_runs_and_reports_endpoints(self) -> None:
        application, base = self._server()
        
        # Add test run
        application.test_history._create(
            job={
                "sheet": "计算器",
                "case_id": "CALC_001",
                "project": "620C_W6830",
                "finished_at": "2026-08-20T10:00:00+08:00",
            },
            result={"verdict": "PASS"},
            stdout="PASS",
            stderr="",
        )
        failure_history_id = application.test_history._create(
            job={
                "sheet": "计算器",
                "case_id": "CALC_003",
                "project": "620C_W6830",
                "finished_at": "2026-08-20T11:00:00+08:00",
            },
            result={"verdict": "ERROR", "reason": "GUI_PING 超时"},
            stdout="",
            stderr="GUI_PING 超时",
        )
        
        # 1. Tests Jobs
        with urlopen(base + "/api/tests/jobs?project=620C_W6830&status=running", timeout=3) as resp:
            data = json.loads(resp.read().decode("utf-8"))
            self.assertEqual(resp.status, 200)
            self.assertIn("items", data)
            self.assertIn("summary", data)
            
        # 2. Reports Summary
        with urlopen(base + "/api/reports/summary?project=620C_W6830", timeout=3) as resp:
            data = json.loads(resp.read().decode("utf-8"))
            self.assertEqual(resp.status, 200)
            self.assertIn("metrics", data)
            self.assertIn("distribution", data)
            self.assertIn("trend", data)
            self.assertIn("top_fail_modules", data)
            self.assertEqual(data["recent_failures"][0]["history_id"], failure_history_id)
            self.assertEqual(data["recent_failures"][0]["case_id"], "CALC_003")
            
        # 3. Reports Runs
        with urlopen(base + "/api/reports/runs?scope=case&project=620C_W6830", timeout=3) as resp:
            data = json.loads(resp.read().decode("utf-8"))
            self.assertEqual(resp.status, 200)
            self.assertIn("items", data)
            
        # 4. Reports Export
        with urlopen(base + "/api/reports/export?project=620C_W6830", timeout=3) as resp:
            self.assertEqual(resp.status, 200)
            self.assertIn("application/vnd.openxmlformats-officedocument", resp.headers.get("Content-Type"))

    def test_report_24h_period_filters_runs_batches_and_export_by_timestamp(self) -> None:
        application, base = self._server()
        now = datetime.now().astimezone()
        within = (now - timedelta(hours=23, minutes=30)).isoformat(timespec="seconds")
        outside = (now - timedelta(hours=24, minutes=30)).isoformat(timespec="seconds")

        for case_id, finished_at in (("CALC_001", within), ("CALC_003", outside)):
            application.test_history._create(
                job={
                    "sheet": "计算器",
                    "case_id": case_id,
                    "project": "620C_W6830",
                    "started_at": finished_at,
                    "finished_at": finished_at,
                },
                result={"verdict": "PASS"},
                stdout="PASS",
                stderr="",
            )

        for batch_id, started_at in (("batch-within-24h", within), ("batch-outside-24h", outside)):
            batch_dir = application.paths.runtime_jobs / batch_id
            batch_dir.mkdir(parents=True)
            (batch_dir / BATCH_STATE_FILE).write_text(
                json.dumps({
                    "project": "620C_W6830",
                    "status": "completed",
                    "started_at": started_at,
                    "finished_at": started_at,
                    "completed": 1,
                    "total": 1,
                }),
                encoding="utf-8",
            )

        query = urlencode({"project": "620C_W6830", "period": "24h"})
        with urlopen(base + f"/api/reports/summary?{query}", timeout=3) as resp:
            summary = json.loads(resp.read().decode("utf-8"))
        self.assertEqual(summary["metrics"]["total"], 1)

        with urlopen(base + f"/api/reports/runs?scope=case&{query}", timeout=3) as resp:
            case_runs = json.loads(resp.read().decode("utf-8"))
        self.assertEqual(case_runs["total"], 1)
        self.assertEqual(case_runs["items"][0]["case_id"], "CALC_001")

        with urlopen(base + f"/api/reports/runs?scope=batch&{query}", timeout=3) as resp:
            batch_runs = json.loads(resp.read().decode("utf-8"))
        self.assertEqual(batch_runs["total"], 1)
        self.assertEqual(batch_runs["items"][0]["batch_id"], "batch-within-24h")

        with urlopen(base + f"/api/reports/export?{query}", timeout=3) as resp:
            import io
            import openpyxl

            workbook = openpyxl.load_workbook(io.BytesIO(resp.read()))
        summary_sheet = workbook["测试报告概览"]
        self.assertEqual(summary_sheet["B2"].value, "最近24小时")
        self.assertEqual(summary_sheet["B3"].value, 1)

    def test_reports_export_includes_each_non_pass_execution(self) -> None:
        application, base = self._server()
        screenshot = Path(self.temporary.name) / "report-failure.bmp"
        screenshot.write_bytes(b"BM-report-failure")

        def create_run(
            *,
            verdict: str,
            case_id: str,
            finished_at: str,
            reason: str,
            batch_id: str = "",
            command: str = "",
            with_screenshot: bool = False,
        ) -> None:
            job = {
                "sheet": "计算器",
                "case_id": case_id,
                "project": "620C_W6830",
                "started_at": finished_at,
                "finished_at": finished_at,
                "case": {
                    "priority": "P0",
                    "precondition_text": "已进入计算器",
                    "steps_text": "点击等号并观察结果",
                    "expected_text": "结果区域显示正确数值",
                },
            }
            if batch_id:
                job["batch_id"] = batch_id
            result = {
                "verdict": verdict,
                "reason": reason,
                "command_trace": ([{
                    "index": 1,
                    "phase": "action",
                    "status": "accepted",
                    "command": command,
                }] if command else []),
                "screenshots": ([{
                    "path": str(screenshot),
                    "label": "失败画面",
                    "phase": "action",
                    "command": command,
                }] if with_screenshot else []),
            }
            application.test_history._create(
                job=job,
                result=result,
                stdout="done",
                stderr="",
            )

        create_run(
            verdict="PASS",
            case_id="CALC_009",
            finished_at="2026-08-20T09:00:00+08:00",
            reason="符合预期",
        )
        create_run(
            verdict="FAIL",
            case_id="CALC_001",
            finished_at="2026-08-20T10:00:00+08:00",
            reason="实际显示 8，与预期显示 9 不一致。",
            batch_id="batch-fail",
            command=":TP_CLICK:10,20,1",
            with_screenshot=True,
        )
        create_run(
            verdict="ERROR",
            case_id="CALC_001",
            finished_at="2026-08-20T11:00:00+08:00",
            reason=(
                "Traceback (most recent call last):\n"
                "  File \"test.py\", line 1, in run\n"
                "HardwareSerialTimeoutError: GUI_PING 超时"
            ),
            command=":GUI_PING:1",
        )
        create_run(
            verdict="CANNOT_VERIFY",
            case_id="CALC_002",
            finished_at="2026-08-20T12:00:00+08:00",
            reason="缺少独立截图证据，无法进行视觉判定。",
            batch_id="batch-cannot-verify",
        )

        query = urlencode({
            "project": "620C_W6830",
            "from": "2026-08-20",
            "to": "2026-08-20",
        })
        with urlopen(base + f"/api/reports/export?{query}", timeout=3) as resp:
            import io
            import openpyxl

            workbook = openpyxl.load_workbook(io.BytesIO(resp.read()))

        self.assertEqual(
            workbook.sheetnames,
            ["测试报告概览", "高频失败模块", "异常用例明细"],
        )
        summary = workbook["测试报告概览"]
        self.assertEqual(summary["B3"].value, 4)
        self.assertEqual(summary["B4"].value, 1)
        self.assertEqual(summary["B5"].value, 1)
        self.assertEqual(summary["B6"].value, 1)
        self.assertEqual(summary["B7"].value, 1)

        detail = workbook["异常用例明细"]
        headers = [cell.value for cell in detail[1]]
        self.assertEqual(headers, [
            "运行时间", "模块", "用例编号", "优先级", "结果", "异常类别",
            "前置条件", "测试步骤", "预期结果", "原因摘要", "原始判定/错误详情",
            "批次编号", "运行记录编号", "执行命令", "截图证据数量",
        ])
        rows = [dict(zip(headers, values)) for values in detail.iter_rows(min_row=2, values_only=True)]
        self.assertEqual([row["结果"] for row in rows], ["CANNOT_VERIFY", "ERROR", "FAIL"])
        self.assertEqual(sum(row["用例编号"] == "CALC_001" for row in rows), 2)

        by_verdict = {row["结果"]: row for row in rows}
        self.assertEqual(by_verdict["FAIL"]["异常类别"], "产品失败")
        self.assertEqual(by_verdict["FAIL"]["批次编号"], "batch-fail")
        self.assertEqual(by_verdict["FAIL"]["截图证据数量"], 1)
        self.assertIn(":TP_CLICK:10,20,1 [action/accepted]", by_verdict["FAIL"]["执行命令"])
        self.assertEqual(
            by_verdict["ERROR"]["原因摘要"],
            "HardwareSerialTimeoutError: GUI_PING 超时",
        )
        self.assertIn("Traceback", by_verdict["ERROR"]["原始判定/错误详情"])
        self.assertEqual(by_verdict["ERROR"]["异常类别"], "执行异常")
        self.assertEqual(by_verdict["CANNOT_VERIFY"]["异常类别"], "无法验证")
        self.assertEqual(by_verdict["CANNOT_VERIFY"]["批次编号"], "batch-cannot-verify")
        self.assertEqual(detail.freeze_panes, "A2")
        self.assertEqual(detail.auto_filter.ref, "A1:O4")

        summary_header = summary["A1"]
        self.assertEqual(summary_header.font.name, "宋体")
        self.assertEqual(summary_header.font.sz, 11)
        self.assertTrue(summary_header.font.bold)
        self.assertEqual(summary_header.border.left.style, "thin")
        self.assertEqual(summary["B1"].font.name, "Times New Roman")
        self.assertEqual(summary["B1"].font.sz, 10)
        self.assertEqual(summary.column_dimensions["A"].width, 22)
        self.assertEqual(summary.page_setup.orientation, "portrait")

        module_sheet = workbook["高频失败模块"]
        self.assertEqual(module_sheet["A1"].font.name, "宋体")
        self.assertEqual(module_sheet["A1"].font.sz, 11)
        self.assertEqual(module_sheet["A1"].border.bottom.style, "thin")
        self.assertEqual(module_sheet.freeze_panes, "A2")

        verdict_rows = {
            detail.cell(row_index, 5).value: row_index
            for row_index in range(2, detail.max_row + 1)
        }
        error_row = verdict_rows["ERROR"]
        self.assertEqual(detail["A1"].font.name, "宋体")
        self.assertEqual(detail["A1"].font.sz, 11)
        self.assertEqual(detail.cell(error_row, 3).font.name, "Times New Roman")
        self.assertEqual(detail.cell(error_row, 5).font.name, "Times New Roman")
        self.assertTrue(detail.cell(error_row, 5).font.bold)
        self.assertTrue(detail.cell(error_row, 11).alignment.wrap_text)
        self.assertEqual(detail.cell(error_row, 11).border.right.style, "thin")
        self.assertEqual(detail.column_dimensions["K"].width, 60)
        self.assertGreater(detail.row_dimensions[error_row].height, 22)
        self.assertLessEqual(detail.row_dimensions[error_row].height, 120)
        self.assertEqual(detail.page_setup.orientation, "landscape")

    def test_environments_and_config_endpoints(self) -> None:
        _, base = self._server()
        
        # 1. Environments list
        with urlopen(base + "/api/environments", timeout=3) as resp:
            data = json.loads(resp.read().decode("utf-8"))
            self.assertEqual(resp.status, 200)
            self.assertIn("items", data)
            self.assertTrue(len(data["items"]) >= 3)
            hardware = next(item for item in data["items"] if item["id"] == "6202_W5230")
            self.assertEqual(hardware["checks"][0]["key"], "profile")
            self.assertEqual(hardware["checks"][0]["label"], "真机运行时档案")
            
        # 2. Environment check
        with self._post_json(base + "/api/environments/620C_W6830/check", {}) as resp:
            data = json.loads(resp.read().decode("utf-8"))
            self.assertEqual(resp.status, 200)
            self.assertIn("result", data)
            
        # 3. Environment PUT
        put_payload = {
            "paths": {
                "source_root": "D:\\Agent-loop-workspace\\620C_W6830",
                "workspace_root": "D:\\Agent-loop-workspace\\620C_W6830",
            }
        }
        with self._put_json(base + "/api/environments/620C_W6830", put_payload) as resp:
            data = json.loads(resp.read().decode("utf-8"))
            self.assertEqual(resp.status, 200)
            self.assertEqual(data["status"], "ok")
            
        # 4. Config GET & POST
        from agent_loop_system.tools.llm_retry import record_actual_llm_success

        record_actual_llm_success(
            root=self.paths.root,
            at="2026-08-21T14:30:00+08:00",
        )
        with urlopen(base + "/api/config", timeout=3) as resp:
            cfg = json.loads(resp.read().decode("utf-8"))
            self.assertEqual(resp.status, 200)
            self.assertIn("llm", cfg)
            self.assertIn("ones", cfg)
            self.assertIn("hardware", cfg)
            self.assertIn("simulator", cfg)
            self.assertEqual(cfg["hardware"]["profile_root"], str(self.paths.root / "profiles"))
            self.assertIn("ble_address", cfg["hardware"])
            self.assertEqual(cfg["hardware"]["ble_scan_timeout"], 15.0)
            self.assertNotIn("hardware_source_root", cfg["simulator"])
            self.assertNotIn("hardware_workspace_root", cfg["simulator"])
            self.assertFalse(cfg["llm"]["configured"])
            self.assertEqual(cfg["llm"]["status"], "unconfigured")
            self.assertEqual(
                cfg["llm"]["last_actual_success_at"],
                "2026-08-21T14:30:00+08:00",
            )

        hardware_profile_root = str(self.paths.root / "portable-profiles")
        with self._put_json(
            base + "/api/environments/6202_W5230",
            {
                "paths": {
                    "profile_root": hardware_profile_root,
                    "profile_version": "v30-test.1",
                }
            },
        ) as resp:
            data = json.loads(resp.read().decode("utf-8"))
            self.assertEqual(resp.status, 200)
            self.assertEqual(data["status"], "ok")
        with urlopen(base + "/api/config", timeout=3) as resp:
            cfg = json.loads(resp.read().decode("utf-8"))
            self.assertEqual(cfg["hardware"]["profile_root"], hardware_profile_root)
            self.assertEqual(cfg["hardware"]["profile_version"], "v30-test.1")
            
        post_cfg = {
            "llm": {"model": "gpt-4o-mini", "timeout": 60},
            "hardware": {
                "capture_provider": "ble",
                "ble_address": "",
                "ble_scan_timeout": 12,
            },
        }
        with self._post_json(base + "/api/config", post_cfg) as resp:
            data = json.loads(resp.read().decode("utf-8"))
            self.assertEqual(resp.status, 200)
            self.assertEqual(data["status"], "ok")
        with urlopen(base + "/api/config", timeout=3) as resp:
            cfg = json.loads(resp.read().decode("utf-8"))
            self.assertEqual(cfg["hardware"]["capture_provider"], "ble")
            self.assertEqual(cfg["hardware"]["ble_address"], "")
            self.assertEqual(cfg["hardware"]["ble_scan_timeout"], 12.0)
            
        # 5. LLM Connectivity check
        with patch("agent_loop_system.tools.llm_config.test_llm_connectivity", return_value={"ok": True, "latency_ms": 120, "model": "gpt-5.6-sol", "message": "连接成功"}):
            with self._post_json(base + "/api/config/test-llm", {}) as resp:
                data = json.loads(resp.read().decode("utf-8"))
                self.assertEqual(resp.status, 200)
                self.assertTrue(data["ok"])
                self.assertEqual(data["model"], "gpt-5.6-sol")

        # 6. Update check & Heartbeat
        with patch(
            "frontend.server.check_for_updates",
            return_value={
                "has_update": False,
                "status": "up_to_date",
                "current_version": "0.4.2",
                "latest_version": "0.4.2",
            },
        ):
            with urlopen(base + "/api/update-check", timeout=3) as resp:
                data = json.loads(resp.read().decode("utf-8"))
                self.assertEqual(resp.status, 200)
                self.assertIn("current_version", data)
            
        with urlopen(base + "/api/system/heartbeat", timeout=3) as resp:
            data = json.loads(resp.read().decode("utf-8"))
            self.assertEqual(resp.status, 200)
            self.assertEqual(data["status"], "ok")

    def test_ble_device_manager_scans_connects_remembers_and_forgets(self) -> None:
        _, base = self._server()
        discovered = [
            SimpleNamespace(
                address="42:74:DC:C8:0A:02",
                name="oraimo Watch Tank N",
                rssi=-48,
            ),
            SimpleNamespace(
                address="11:22:33:44:55:66",
                name="oraimo Watch Tank N Pro",
                rssi=-71,
            ),
        ]
        scan_mock = AsyncMock(return_value=discovered)
        with (
            patch("frontend.server.scan_watches", scan_mock),
            patch("frontend.server.WatchBleClient") as client_type,
        ):
            query = urlencode({"q": "c8:0a", "timeout": "3"})
            with urlopen(base + f"/api/hardware/ble/devices?{query}", timeout=3) as resp:
                data = json.loads(resp.read().decode("utf-8"))
            self.assertEqual(resp.status, 200)
            self.assertEqual(len(data["items"]), 1)
            self.assertEqual(data["items"][0]["address"], "42:74:DC:C8:0A:02")
            self.assertEqual(data["items"][0]["status"], "discovered")
            self.assertFalse(data["items"][0]["connected"])
            scan_mock.assert_awaited_once_with(timeout=3.0, address=None)
            client_type.assert_not_called()

        fake_client = SimpleNamespace(
            connect=AsyncMock(return_value=None),
            close=AsyncMock(return_value=None),
            connected=True,
        )
        with patch("frontend.server.WatchBleClient", return_value=fake_client) as client_type:
            with self._post_json(
                base + "/api/hardware/ble/connect",
                {
                    "address": "42:74:DC:C8:0A:02",
                    "name": "oraimo Watch Tank N",
                    "timeout": 4,
                },
            ) as resp:
                data = json.loads(resp.read().decode("utf-8"))
            self.assertEqual(resp.status, 200)
            self.assertTrue(data["verified"])
            self.assertFalse(data["connected"])
            self.assertEqual(data["connection_mode"], "on_demand")
            self.assertEqual(data["device"]["status"], "verified")
            client_type.assert_called_once()
            fake_client.connect.assert_awaited_once_with(pair=False)
            fake_client.close.assert_awaited_once_with()

        remembered_path = self.paths.root / ".runtime" / "ble-devices.json"
        remembered_payload = json.loads(remembered_path.read_text(encoding="utf-8"))
        self.assertEqual(remembered_payload["schema_version"], 1)
        self.assertEqual(len(remembered_payload["items"]), 1)
        self.assertEqual(
            remembered_payload["items"][0]["address"],
            "42:74:DC:C8:0A:02",
        )
        self.assertEqual(
            os.environ["W30_HARDWARE_BLE_ADDRESS"],
            "42:74:DC:C8:0A:02",
        )

        with (
            patch("frontend.server.scan_watches", new=AsyncMock()) as scan_again,
            patch("frontend.server.WatchBleClient") as client_again,
        ):
            with urlopen(base + "/api/hardware/ble/remembered", timeout=3) as resp:
                remembered = json.loads(resp.read().decode("utf-8"))
            self.assertEqual(len(remembered["items"]), 1)
            self.assertTrue(remembered["items"][0]["selected"])
            self.assertEqual(remembered["items"][0]["status"], "verified")
            self.assertFalse(remembered["items"][0]["connected"])
            scan_again.assert_not_awaited()
            client_again.assert_not_called()

        failed_client = SimpleNamespace(
            connect=AsyncMock(side_effect=RuntimeError("device unavailable")),
            close=AsyncMock(return_value=None),
            connected=False,
        )
        with patch("frontend.server.WatchBleClient", return_value=failed_client):
            with self.assertRaises(HTTPError) as context:
                self._post_json(
                    base + "/api/hardware/ble/connect",
                    {
                        "address": "AA:BB:CC:DD:EE:FF",
                        "name": "unavailable watch",
                        "timeout": 2,
                    },
                )
            self.assertEqual(context.exception.code, 502)
            error = json.loads(context.exception.read().decode("utf-8"))
            self.assertFalse(error["verified"])
            failed_client.close.assert_awaited_once_with()
        remembered_payload = json.loads(remembered_path.read_text(encoding="utf-8"))
        self.assertEqual(len(remembered_payload["items"]), 1)

        address = quote("42:74:DC:C8:0A:02", safe="")
        with (
            patch("frontend.server.scan_watches", new=AsyncMock()) as passive_scan,
            patch("frontend.server.WatchBleClient") as passive_client,
        ):
            with self._delete_json(base + f"/api/hardware/ble/remembered/{address}") as resp:
                deleted = json.loads(resp.read().decode("utf-8"))
            self.assertTrue(deleted["deleted"])
            self.assertEqual(deleted["selected_address"], "")
            passive_scan.assert_not_awaited()
            passive_client.assert_not_called()
        self.assertEqual(os.environ["W30_HARDWARE_BLE_ADDRESS"], "")
        with urlopen(base + "/api/hardware/ble/remembered", timeout=3) as resp:
            remembered = json.loads(resp.read().decode("utf-8"))
        self.assertEqual(remembered["items"], [])

    def test_hardware_serial_ports_endpoint(self) -> None:
        _, base = self._server()
        os.environ["W30_HARDWARE_PORT"] = "COM7"

        mock_payload = {
            "configured_port": "COM7",
            "selected_port": "COM7",
            "default_port": "COM7",
            "active_count": 1,
            "items": [
                {
                    "port": "COM7",
                    "friendly_name": "USB Serial Port (COM7)",
                    "description": "USB Serial Port",
                    "hardware_id": r"FTDIBUS\VID_0403+PID_6001\0000",
                    "manufacturer": "FTDI",
                    "kind": "usb",
                    "kind_label": "USB 串口",
                    "present": True,
                    "supercom_open": True,
                    "pipe_name": "SuperCom.AgentBridge.COM7",
                    "pipe_path": r"\\.\pipe\SuperCom.AgentBridge.COM7",
                    "missing": False,
                },
                {
                    "port": "COM1",
                    "friendly_name": "Communications Port (COM1)",
                    "description": "Communications Port",
                    "hardware_id": r"ACPI\PNP0501\1",
                    "manufacturer": "Standard",
                    "kind": "system",
                    "kind_label": "系统/板载串口",
                    "present": True,
                    "supercom_open": False,
                    "pipe_name": "SuperCom.AgentBridge.COM1",
                    "pipe_path": r"\\.\pipe\SuperCom.AgentBridge.COM1",
                    "missing": False,
                },
            ],
            "available": True,
        }

        with patch(
            "agent_loop_system.tools.hardware_serial_ports.get_serial_ports_status",
            return_value=mock_payload,
        ) as mock_get_status:
            with urlopen(base + "/api/hardware/serial-ports", timeout=3) as resp:
                data = json.loads(resp.read().decode("utf-8"))
            self.assertEqual(resp.status, 200)
            self.assertEqual(data["configured_port"], "COM7")
            self.assertEqual(data["selected_port"], "COM7")
            self.assertEqual(data["active_count"], 1)
            self.assertEqual(len(data["items"]), 2)
            self.assertTrue(data["items"][0]["supercom_open"])
            self.assertEqual(data["items"][0]["kind"], "usb")
            self.assertFalse(data["items"][1]["supercom_open"])
            self.assertEqual(data["items"][1]["kind"], "system")
            mock_get_status.assert_called_once_with("COM7")

        post_cfg = {
            "hardware": {
                "port": "COM8",
                "baudrate": 1500000,
                "transport": "supercom",
                "capture_provider": "mtp",
            }
        }
        with self._post_json(base + "/api/config", post_cfg) as resp:
            data = json.loads(resp.read().decode("utf-8"))
            self.assertEqual(resp.status, 200)
            self.assertEqual(data["status"], "ok")
        self.assertEqual(os.environ["W30_HARDWARE_PORT"], "COM8")


if __name__ == "__main__":
    unittest.main()
