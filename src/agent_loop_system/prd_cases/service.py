"""Persistent PRD -> Excel -> review -> case-management workflow.

The bundled QA Skill is treated as a pinned, read-only rule package.  Generated
rows are always reopened from the exported workbook before they are exposed to
reviewers; approvals and sync operations are consequently bound to the exact
workbook SHA-256 rather than to an in-memory draft.
"""
from __future__ import annotations

import base64
import hashlib
import json
import re
import threading
import uuid
import zipfile
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from typing import Any, Callable

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, Field

from agent_loop_system.tools.llm_config import create_chat_llm, get_llm_api_key


QA_SKILL_FILENAME = "xiaozhou-portable-skill-execution-quality-20260825.zip"
QA_SKILL_SHA256 = "FFDF484A2F24F30A7A9CAD757F8545F4ECEC2F990B5DE9768A6BB9EA93DB6DBA"
SUPPORTED_PRD_SUFFIXES = {".md", ".txt", ".docx"}
MAX_PRD_BYTES = 20 * 1024 * 1024
CASE_ID_RE = re.compile(r"^[A-Za-z0-9_.-]+$")
EXCEL_HEADERS = [
    "用例编号", "功能模块", "功能点", "测试项", "测试点", "用例标题", "优先级",
    "前置条件", "操作步骤", "预期结果", "测试类型", "需求ID", "备注", "稳定ID",
]


def _now() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def _atomic_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def _safe_excel_text(value: Any) -> str:
    text = str(value or "").strip()
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


class GeneratedCase(BaseModel):
    case_id: str = Field(description="稳定且唯一的英文、数字、点、下划线或连字符用例编号")
    functional_module: str
    feature: str
    test_item: str
    test_point: str
    title: str
    priority: str = "P1"
    preconditions: str
    steps: list[str]
    expected_results: list[str]
    test_type: str = "功能"
    requirement_ids: list[str] = Field(default_factory=list)
    note: str = ""


class GeneratedCaseBatch(BaseModel):
    cases: list[GeneratedCase]


Generator = Callable[[str, dict[str, Any]], list[dict[str, Any]]]


class PrdCaseService:
    """Owns asynchronous generation jobs and their immutable review evidence."""

    def __init__(
        self,
        root: Path,
        case_repository: Any,
        *,
        project_lookup: Callable[[str], dict[str, Any]],
        generator: Generator | None = None,
        skill_bundle: Path | None = None,
        start_threads: bool = True,
    ) -> None:
        self.root = Path(root)
        self.jobs_root = self.root / "project_data" / "prd_cases" / "jobs"
        self.jobs_root.mkdir(parents=True, exist_ok=True)
        self.case_repository = case_repository
        self.project_lookup = project_lookup
        self.generator = generator or self._llm_generate
        self.skill_bundle = Path(skill_bundle or (
            self.root / "resources" / "skills" / QA_SKILL_FILENAME
        ))
        self.start_threads = start_threads
        self._lock = threading.RLock()
        self._threads: dict[str, threading.Thread] = {}

    def _job_dir(self, job_id: str) -> Path:
        if not re.fullmatch(r"prd-[a-f0-9]{32}", str(job_id)):
            raise ValueError("PRD_JOB_ID_INVALID")
        return self.jobs_root / job_id

    def _state_path(self, job_id: str) -> Path:
        return self._job_dir(job_id) / "state.json"

    def _read_state(self, job_id: str) -> dict[str, Any]:
        path = self._state_path(job_id)
        if not path.is_file():
            raise ValueError("PRD_JOB_NOT_FOUND")
        try:
            value = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, UnicodeError, json.JSONDecodeError) as exc:
            raise RuntimeError("PRD_JOB_STATE_CORRUPT") from exc
        if not isinstance(value, dict):
            raise RuntimeError("PRD_JOB_STATE_CORRUPT")
        return value

    def _write_state(self, state: dict[str, Any]) -> None:
        state["updated_at"] = _now()
        _atomic_json(self._state_path(str(state["job_id"])), state)

    def _append_audit(self, job_id: str, event: str, details: dict[str, Any]) -> None:
        path = self._job_dir(job_id) / "audit.jsonl"
        record = {"at": _now(), "event": event, "details": details}
        with path.open("a", encoding="utf-8") as handle:
            handle.write(json.dumps(record, ensure_ascii=False, separators=(",", ":")) + "\n")

    def verify_skill_bundle(self) -> dict[str, Any]:
        if not self.skill_bundle.is_file():
            raise RuntimeError("QA_SKILL_MISSING: 便携 QA Skill 未随程序发布")
        actual = _sha256(self.skill_bundle)
        if actual != QA_SKILL_SHA256:
            raise RuntimeError(
                f"QA_SKILL_HASH_MISMATCH: expected={QA_SKILL_SHA256} actual={actual}"
            )
        try:
            with zipfile.ZipFile(self.skill_bundle) as archive:
                names = set(archive.namelist())
                required = {
                    "xiaozhou/qa/workflow-v2/SKILL.md",
                    "xiaozhou/qa/test-planning/SKILL.md",
                    "xiaozhou/qa/workflow-v2/scripts/release_acceptance.py",
                }
                missing = sorted(required - names)
        except (OSError, zipfile.BadZipFile) as exc:
            raise RuntimeError("QA_SKILL_ZIP_INVALID") from exc
        if missing:
            raise RuntimeError("QA_SKILL_CONTENT_MISSING: " + ", ".join(missing))
        return {"filename": self.skill_bundle.name, "sha256": actual, "verified": True}

    def create_job(
        self,
        *,
        project_id: str,
        filename: str,
        file_base64: str,
        execution_profile: str = "core",
    ) -> dict[str, Any]:
        project = self.project_lookup(str(project_id).strip())
        safe_name = Path(str(filename or "")).name
        suffix = Path(safe_name).suffix.lower()
        if suffix not in SUPPORTED_PRD_SUFFIXES:
            raise ValueError("PRD_FILE_TYPE_UNSUPPORTED: 仅支持 .md、.txt、.docx")
        try:
            data = base64.b64decode(file_base64, validate=True)
        except (ValueError, TypeError) as exc:
            raise ValueError("PRD_BASE64_INVALID") from exc
        if not data or len(data) > MAX_PRD_BYTES:
            raise ValueError("PRD_FILE_SIZE_INVALID: 文件必须大于 0 且不超过 20 MB")
        profile = str(execution_profile or "core").strip().lower()
        if profile not in {"core", "strict"}:
            raise ValueError("PRD_EXECUTION_PROFILE_INVALID")

        job_id = "prd-" + uuid.uuid4().hex
        job_dir = self._job_dir(job_id)
        job_dir.mkdir(parents=True, exist_ok=False)
        source_path = job_dir / ("source" + suffix)
        source_path.write_bytes(data)
        state = {
            "schema_version": 1,
            "job_id": job_id,
            "project_id": str(project["project_id"]),
            "filename": safe_name,
            "source_sha256": hashlib.sha256(data).hexdigest().upper(),
            "execution_profile": profile,
            "status": "UPLOADED",
            "progress": 0,
            "stage": "等待生成",
            "case_count": 0,
            "review": None,
            "release": {"decision": "NO_GO", "readiness": "DRAFT"},
            "sync_result": None,
            "created_at": _now(),
            "updated_at": _now(),
        }
        self._write_state(state)
        self._append_audit(job_id, "JOB_CREATED", {
            "project_id": state["project_id"], "source_sha256": state["source_sha256"]
        })
        self.start(job_id)
        return self.get_job(job_id)

    def start(self, job_id: str, *, feedback: str = "") -> None:
        if not self.start_threads:
            return
        with self._lock:
            current = self._threads.get(job_id)
            if current and current.is_alive():
                raise RuntimeError("PRD_JOB_ALREADY_RUNNING")
            thread = threading.Thread(
                target=self.run_job, args=(job_id,), kwargs={"feedback": feedback},
                name=f"prd-cases-{job_id[-8:]}", daemon=True,
            )
            self._threads[job_id] = thread
            thread.start()

    def run_job(self, job_id: str, *, feedback: str = "") -> None:
        try:
            with self._lock:
                state = self._read_state(job_id)
                state.update({"status": "GENERATING", "progress": 5, "stage": "校验 QA Skill"})
                state["review"] = None
                state["sync_result"] = None
                state["release"] = {"decision": "NO_GO", "readiness": "DRAFT"}
                self._write_state(state)
            skill = self.verify_skill_bundle()
            state["skill"] = skill
            state.update({"progress": 15, "stage": "解析 PRD"})
            self._write_state(state)
            source_path = next(self._job_dir(job_id).glob("source.*"), None)
            if source_path is None:
                raise RuntimeError("PRD_SOURCE_MISSING")
            text = self._parse_prd(source_path)
            (self._job_dir(job_id) / "normalized-prd.md").write_text(text, encoding="utf-8")
            state.update({"progress": 30, "stage": "按 QA Skill 设计用例"})
            self._write_state(state)
            context = {
                "job_id": job_id,
                "project_id": state["project_id"],
                "execution_profile": state["execution_profile"],
                "feedback": str(feedback or "").strip(),
                "skill_sha256": skill["sha256"],
            }
            cases = self._normalize_cases(self.generator(text, context))
            state.update({"progress": 72, "stage": "生成并回读 Excel"})
            self._write_state(state)
            workbook_path = self._job_dir(job_id) / "generated-test-cases.xlsx"
            self._export_excel(workbook_path, cases, state)
            reopened = self._read_excel(workbook_path)
            gate = self._quality_gate(reopened)
            workbook_sha = _sha256(workbook_path)
            snapshot = {
                "schema_version": 1,
                "job_id": job_id,
                "workbook_sha256": workbook_sha,
                "skill_sha256": skill["sha256"],
                "generated_at": _now(),
                "cases": reopened,
                "gate": gate,
            }
            _atomic_json(self._job_dir(job_id) / "review-snapshot.json", snapshot)
            state.update({
                "status": "READY_FOR_REVIEW" if gate["passed"] else "QUALITY_BLOCKED",
                "progress": 100,
                "stage": "等待人工审查" if gate["passed"] else "质量门禁未通过",
                "case_count": len(reopened),
                "workbook_sha256": workbook_sha,
                "quality_gate": gate,
                "release": {"decision": "NO_GO", "readiness": "READY_FOR_REVIEW" if gate["passed"] else "DRAFT"},
            })
            self._write_state(state)
            self._append_audit(job_id, "GENERATION_COMPLETED", {
                "case_count": len(reopened), "workbook_sha256": workbook_sha,
                "gate_passed": gate["passed"], "feedback": context["feedback"],
            })
        except BaseException as exc:
            with self._lock:
                state = self._read_state(job_id)
                state.update({
                    "status": "FAILED", "stage": "生成失败", "error": str(exc),
                    "release": {"decision": "NO_GO", "readiness": "DRAFT"},
                })
                self._write_state(state)
                self._append_audit(job_id, "GENERATION_FAILED", {"error": str(exc)})

    def regenerate(self, job_id: str, feedback: str) -> dict[str, Any]:
        state = self._read_state(job_id)
        if state["status"] not in {"REJECTED", "QUALITY_BLOCKED", "FAILED"}:
            raise ValueError("PRD_REGENERATE_NOT_ALLOWED")
        instruction = str(feedback or "").strip()
        if not instruction:
            raise ValueError("PRD_REGENERATE_FEEDBACK_REQUIRED")
        self._append_audit(job_id, "REGENERATION_REQUESTED", {"feedback": instruction})
        if self.start_threads:
            self.start(job_id, feedback=instruction)
        else:
            self.run_job(job_id, feedback=instruction)
        return self.get_job(job_id)

    def _parse_prd(self, path: Path) -> str:
        if path.suffix.lower() in {".md", ".txt"}:
            try:
                text = path.read_text(encoding="utf-8-sig")
            except UnicodeDecodeError:
                text = path.read_text(encoding="gb18030")
        elif path.suffix.lower() == ".docx":
            try:
                with zipfile.ZipFile(path) as archive:
                    xml = archive.read("word/document.xml").decode("utf-8")
            except (OSError, KeyError, UnicodeError, zipfile.BadZipFile) as exc:
                raise ValueError("PRD_DOCX_INVALID") from exc
            xml = re.sub(r"</w:p>", "\n", xml)
            xml = re.sub(r"<w:tab[^>]*/>", "\t", xml)
            text = re.sub(r"<[^>]+>", "", xml)
            text = (text.replace("&lt;", "<").replace("&gt;", ">").replace("&amp;", "&"))
        else:
            raise ValueError("PRD_FILE_TYPE_UNSUPPORTED")
        text = text.replace("\x00", "").strip()
        if len(text) < 20:
            raise ValueError("PRD_CONTENT_TOO_SHORT: 未解析到足够的需求内容")
        if len(text) > 180_000:
            raise ValueError("PRD_CONTENT_TOO_LONG: 当前版本最多处理 18 万字符")
        return text

    def _skill_prompt_excerpt(self) -> str:
        with zipfile.ZipFile(self.skill_bundle) as archive:
            parts = []
            for name in (
                "xiaozhou/qa/workflow-v2/SKILL.md",
                "xiaozhou/qa/test-planning/SKILL.md",
            ):
                parts.append(archive.read(name).decode("utf-8")[:9000])
        return "\n\n".join(parts)

    def _llm_generate(self, prd_text: str, context: dict[str, Any]) -> list[dict[str, Any]]:
        if not get_llm_api_key("exploration"):
            raise RuntimeError("LLM_AUTH_FAILED: 未配置用于 PRD 转用例的大模型 API Key")
        llm = create_chat_llm(api_key_scope="exploration")
        if llm is None:
            raise RuntimeError("LLM_DEPENDENCY_MISSING")
        rules = self._skill_prompt_excerpt()
        sections = self._split_prd(prd_text)
        generated: list[dict[str, Any]] = []
        for index, section in enumerate(sections, start=1):
            prompt = f"""你是 QA 测试用例设计执行器。必须遵循下方已校验 xiaozhou QA Skill 的规则摘要，
尤其严格使用 功能模块 > 功能点 > 测试项 > 测试点 > 用例详情 的层级；按用户可观察行为设计原子用例，
步骤有明确对象与顺序，预期可验证，禁止只写“正常/正确”，不得臆造需求外能力。

Skill SHA256: {context['skill_sha256']}
执行画像: {context['execution_profile']}
目标项目: {context['project_id']}
分段: {index}/{len(sections)}
审查修改意见: {context.get('feedback') or '无'}

Skill 规则摘要：
{rules}

PRD 内容：
{section}

请返回结构化用例。case_id 必须稳定、唯一，仅使用英文字母、数字、点、下划线或连字符；
同一业务的正常、异常、边界、状态迁移分别拆分。每个步骤和预期都写成独立数组项。"""
            try:
                result = llm.with_structured_output(GeneratedCaseBatch).invoke(prompt)
            except Exception as exc:
                raise RuntimeError(f"PRD_LLM_GENERATION_FAILED: {exc}") from exc
            batch = result if isinstance(result, GeneratedCaseBatch) else GeneratedCaseBatch.model_validate(result)
            generated.extend(item.model_dump() for item in batch.cases)
        return generated

    @staticmethod
    def _split_prd(text: str, limit: int = 28_000) -> list[str]:
        if len(text) <= limit:
            return [text]
        chunks: list[str] = []
        current = ""
        for block in re.split(r"(?=\n#{1,3}\s+)", text):
            if current and len(current) + len(block) > limit:
                chunks.append(current.strip())
                current = block
            else:
                current += block
        if current.strip():
            chunks.append(current.strip())
        return chunks or [text]

    def _normalize_cases(self, raw_cases: Any) -> list[dict[str, Any]]:
        if not isinstance(raw_cases, list) or not raw_cases:
            raise ValueError("PRD_NO_CASES_GENERATED")
        result: list[dict[str, Any]] = []
        seen: set[str] = set()
        for index, raw in enumerate(raw_cases, start=1):
            item = GeneratedCase.model_validate(raw).model_dump()
            case_id = str(item["case_id"]).strip()
            if not CASE_ID_RE.fullmatch(case_id):
                raise ValueError(f"PRD_CASE_ID_INVALID: row={index} id={case_id}")
            if case_id in seen:
                raise ValueError(f"PRD_CASE_ID_DUPLICATE: {case_id}")
            seen.add(case_id)
            item["case_id"] = case_id
            item["priority"] = str(item.get("priority") or "P1").upper()
            item["stable_id"] = case_id
            result.append(item)
        return result

    def _export_excel(self, path: Path, cases: list[dict[str, Any]], state: dict[str, Any]) -> None:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "测试用例"
        sheet.append(EXCEL_HEADERS)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="24445C")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for case in cases:
            steps = "\n".join(f"{i}. {_safe_excel_text(value)}" for i, value in enumerate(case["steps"], 1))
            expected = "\n".join(f"{i}. {_safe_excel_text(value)}" for i, value in enumerate(case["expected_results"], 1))
            sheet.append([
                _safe_excel_text(case["case_id"]), _safe_excel_text(case["functional_module"]),
                _safe_excel_text(case["feature"]), _safe_excel_text(case["test_item"]),
                _safe_excel_text(case["test_point"]), _safe_excel_text(case["title"]),
                _safe_excel_text(case["priority"]), _safe_excel_text(case["preconditions"]),
                steps, expected, _safe_excel_text(case["test_type"]),
                ", ".join(_safe_excel_text(value) for value in case["requirement_ids"]),
                _safe_excel_text(case["note"]), _safe_excel_text(case["stable_id"]),
            ])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        widths = [18, 18, 20, 22, 28, 32, 10, 32, 48, 48, 12, 18, 28, 18]
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        meta = workbook.create_sheet("生成信息")
        meta.append(["字段", "值"])
        meta.append(["任务ID", state["job_id"]])
        meta.append(["目标项目", state["project_id"]])
        meta.append(["PRD SHA256", state["source_sha256"]])
        meta.append(["QA Skill SHA256", state["skill"]["sha256"]])
        meta.append(["执行画像", state["execution_profile"]])
        meta.sheet_state = "hidden"
        workbook.save(path)

    def _read_excel(self, path: Path) -> list[dict[str, Any]]:
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=False)
        try:
            sheet = workbook["测试用例"]
            header = [str(cell.value or "").strip() for cell in sheet[1]]
            if header[:len(EXCEL_HEADERS)] != EXCEL_HEADERS:
                raise RuntimeError("PRD_EXCEL_HEADER_MISMATCH")
            rows: list[dict[str, Any]] = []
            for values in sheet.iter_rows(min_row=2, values_only=True):
                if not any(value is not None and str(value).strip() for value in values):
                    continue
                row = dict(zip(EXCEL_HEADERS, values))
                rows.append({
                    "case_id": str(row["用例编号"] or "").lstrip("'").strip(),
                    "functional_module": str(row["功能模块"] or "").strip(),
                    "feature": str(row["功能点"] or "").strip(),
                    "test_item": str(row["测试项"] or "").strip(),
                    "test_point": str(row["测试点"] or "").strip(),
                    "title": str(row["用例标题"] or "").strip(),
                    "priority": str(row["优先级"] or "P1").strip(),
                    "preconditions": str(row["前置条件"] or "").strip(),
                    "steps_text": str(row["操作步骤"] or "").strip(),
                    "expected_text": str(row["预期结果"] or "").strip(),
                    "test_type": str(row["测试类型"] or "").strip(),
                    "requirement_ids": [v.strip() for v in str(row["需求ID"] or "").split(",") if v.strip()],
                    "note": str(row["备注"] or "").strip(),
                    "stable_id": str(row["稳定ID"] or "").lstrip("'").strip(),
                })
            return rows
        finally:
            workbook.close()

    @staticmethod
    def _quality_gate(cases: list[dict[str, Any]]) -> dict[str, Any]:
        errors: list[dict[str, str]] = []
        warnings: list[dict[str, str]] = []
        seen: set[str] = set()
        for case in cases:
            case_id = case.get("case_id", "")
            for field in ("functional_module", "feature", "test_item", "test_point", "title", "steps_text", "expected_text"):
                if not str(case.get(field) or "").strip():
                    errors.append({"case_id": case_id, "code": "REQUIRED_FIELD_EMPTY", "message": f"{field} 不能为空"})
            if case_id in seen or not CASE_ID_RE.fullmatch(str(case_id)):
                errors.append({"case_id": case_id, "code": "STABLE_ID_INVALID", "message": "稳定 ID 重复或格式不合法"})
            seen.add(str(case_id))
            if case.get("stable_id") != case_id:
                errors.append({"case_id": case_id, "code": "STABLE_ID_MISMATCH", "message": "稳定 ID 与用例编号不一致"})
            expected = str(case.get("expected_text") or "")
            if re.fullmatch(r"(?:\d+[.、]\s*)?(正常|正确|符合预期)[。.]?", expected):
                errors.append({"case_id": case_id, "code": "EXPECTED_NOT_OBSERVABLE", "message": "预期结果不可只写正常/正确"})
            if len(str(case.get("steps_text") or "")) < 6:
                errors.append({"case_id": case_id, "code": "STEP_TOO_VAGUE", "message": "操作步骤过于笼统"})
            if case.get("functional_module") == case.get("feature"):
                warnings.append({"case_id": case_id, "code": "CLASSIFICATION_REVIEW", "message": "功能模块与功能点相同，请人工确认分类"})
        return {
            "passed": bool(cases) and not errors,
            "checks": {
                "wording": "PASS" if not any(e["code"] in {"EXPECTED_NOT_OBSERVABLE", "STEP_TOO_VAGUE"} for e in errors) else "FAIL",
                "classification": "PASS" if not any(w["code"] == "CLASSIFICATION_REVIEW" for w in warnings) else "WARN",
                "order": "PASS" if not any(e["code"] == "STEP_TOO_VAGUE" for e in errors) else "FAIL",
                "feature_boundary": "PASS" if cases else "FAIL",
                "stable_ids": "PASS" if not any(e["code"].startswith("STABLE_ID") for e in errors) else "FAIL",
            },
            "errors": errors,
            "warnings": warnings,
        }

    def list_jobs(self, *, project_id: str = "") -> list[dict[str, Any]]:
        items: list[dict[str, Any]] = []
        for path in self.jobs_root.glob("prd-*/state.json"):
            try:
                state = json.loads(path.read_text(encoding="utf-8"))
            except (OSError, UnicodeError, json.JSONDecodeError):
                continue
            if project_id and state.get("project_id") != project_id:
                continue
            items.append(self._public_state(state))
        return sorted(items, key=lambda item: str(item.get("created_at", "")), reverse=True)

    @staticmethod
    def _public_state(state: dict[str, Any]) -> dict[str, Any]:
        return deepcopy({key: value for key, value in state.items() if key not in {"internal_path"}})

    def get_job(self, job_id: str) -> dict[str, Any]:
        return self._public_state(self._read_state(job_id))

    def get_cases(self, job_id: str) -> dict[str, Any]:
        snapshot_path = self._job_dir(job_id) / "review-snapshot.json"
        if not snapshot_path.is_file():
            return {"job_id": job_id, "items": [], "gate": None}
        snapshot = json.loads(snapshot_path.read_text(encoding="utf-8"))
        return {
            "job_id": job_id,
            "workbook_sha256": snapshot["workbook_sha256"],
            "items": snapshot["cases"],
            "gate": snapshot["gate"],
        }

    def workbook_path(self, job_id: str) -> Path:
        path = self._job_dir(job_id) / "generated-test-cases.xlsx"
        if not path.is_file():
            raise ValueError("PRD_WORKBOOK_NOT_READY")
        return path

    def review(
        self, job_id: str, *, action: str, reviewer: str, comment: str = "",
        case_comments: list[dict[str, Any]] | None = None,
    ) -> dict[str, Any]:
        with self._lock:
            state = self._read_state(job_id)
            if state["status"] != "READY_FOR_REVIEW":
                raise ValueError("PRD_REVIEW_NOT_ALLOWED")
            reviewer = str(reviewer or "").strip()
            if not reviewer:
                raise ValueError("PRD_REVIEWER_REQUIRED")
            action = str(action or "").strip().lower()
            if action not in {"approve", "reject"}:
                raise ValueError("PRD_REVIEW_ACTION_INVALID")
            workbook_sha = _sha256(self.workbook_path(job_id))
            if workbook_sha != state.get("workbook_sha256"):
                raise ValueError("PRD_WORKBOOK_CHANGED: Excel 已变化，请重新生成审查快照")
            comments = case_comments if isinstance(case_comments, list) else []
            review = {
                "decision": "APPROVED" if action == "approve" else "REJECTED",
                "reviewer": reviewer,
                "comment": str(comment or "").strip(),
                "case_comments": comments,
                "workbook_sha256": workbook_sha,
                "reviewed_at": _now(),
            }
            if action == "reject" and not review["comment"] and not comments:
                raise ValueError("PRD_REJECTION_REASON_REQUIRED")
            state["review"] = review
            state["status"] = "APPROVED" if action == "approve" else "REJECTED"
            state["stage"] = "审查通过，可同步" if action == "approve" else "已驳回，等待修改"
            state["release"] = (
                {"decision": "GO", "readiness": "READY_WITH_RISKS", "workbook_sha256": workbook_sha}
                if action == "approve"
                else {"decision": "NO_GO", "readiness": "DRAFT", "workbook_sha256": workbook_sha}
            )
            self._write_state(state)
            self._append_audit(job_id, "REVIEW_" + action.upper(), review)
        return self.get_job(job_id)

    def sync(self, job_id: str) -> dict[str, Any]:
        with self._lock:
            state = self._read_state(job_id)
            if state["status"] in {"SYNCED", "SYNCED_WITH_CONFLICTS"} and isinstance(state.get("sync_result"), dict):
                return deepcopy(state["sync_result"])
            if state["status"] != "APPROVED" or state.get("release", {}).get("decision") != "GO":
                raise ValueError("PRD_SYNC_REQUIRES_APPROVAL")
            workbook_sha = _sha256(self.workbook_path(job_id))
            if workbook_sha != state.get("workbook_sha256") or workbook_sha != state.get("review", {}).get("workbook_sha256"):
                raise ValueError("PRD_WORKBOOK_CHANGED")
            snapshot = self.get_cases(job_id)
            project = self.project_lookup(state["project_id"])
            cases = [self._to_managed_case(case, project, state) for case in snapshot["items"]]
            result = self.case_repository.sync_approved_prd_cases(
                project,
                cases,
                source_ref={
                    "type": "PRD_APPROVAL",
                    "job_id": job_id,
                    "source_filename": state["filename"],
                    "source_sha256": state["source_sha256"],
                    "workbook_sha256": workbook_sha,
                    "skill_sha256": state["skill"]["sha256"],
                    "reviewer": state["review"]["reviewer"],
                    "reviewed_at": state["review"]["reviewed_at"],
                },
                source_sha256=workbook_sha,
            )
            sync_result = {**result, "job_id": job_id, "project_id": state["project_id"], "synced_at": _now()}
            state["sync_result"] = sync_result
            state["status"] = "SYNCED_WITH_CONFLICTS" if result["conflicts"] else "SYNCED"
            state["stage"] = (
                f"已同步到用例管理，{result['conflicts']} 条编号冲突未写入"
                if result["conflicts"]
                else "已同步到用例管理"
            )
            self._write_state(state)
            self._append_audit(job_id, "SYNC_COMPLETED", sync_result)
            return sync_result

    @staticmethod
    def _to_managed_case(case: dict[str, Any], project: dict[str, Any], state: dict[str, Any]) -> dict[str, Any]:
        return {
            "case_id": case["case_id"],
            "sheet": case["functional_module"],
            "functional_module": case["functional_module"],
            "feature": case["feature"],
            "test_item": case["test_item"],
            "test_point": case["test_point"],
            "title": case["title"],
            "priority": case["priority"],
            "precondition_text": case["preconditions"],
            "steps_text": case["steps_text"],
            "expected_text": case["expected_text"],
            "test_type": case["test_type"],
            "requirement_ids": case["requirement_ids"],
            "note": case["note"],
            "stable_id": case["stable_id"],
            "workflow_state": "ACTIVE",
            "applicable_platforms": list(project.get("allowed_platforms") or []),
            "automation_maturity": "UNMAPPED",
            "mapping_status": "",
            "verification_points": [line for line in case["expected_text"].splitlines() if line.strip()],
            "source_ref": {"job_id": state["job_id"], "workbook_sha256": state["workbook_sha256"]},
        }
